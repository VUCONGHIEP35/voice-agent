# -*- coding: utf-8 -*-
import json, sys
sys.path.insert(0, ".")
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
import data_def as D
import taxonomy_def as TX

OUT_XLSX = "../out/du_lieu_mau_tts_tro_ly_o_to.xlsx"
OUT_JSONL = "../out/du_lieu_mau_tts_tro_ly_o_to.jsonl"

F = "Arial"
f_norm = Font(name=F, size=10)
f_bold = Font(name=F, size=10, bold=True)
f_title = Font(name=F, size=14, bold=True)
f_head = Font(name=F, size=10, bold=True, color="FFFFFF")
f_input = Font(name=F, size=10, color="0000FF")
fill_req = PatternFill("solid", fgColor="1F3864")
fill_opt = PatternFill("solid", fgColor="7F7F7F")
fill_sec = PatternFill("solid", fgColor="D9E1F2")
fill_input = PatternFill("solid", fgColor="FFFF00")
fill_prov = PatternFill("solid", fgColor="548235")
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")
center_wrap = Alignment(wrap_text=True, vertical="center", horizontal="center")


def style_range(ws, ref_rows, ncols, header_fills=None):
    pass


def write_table(ws, start_row, headers, rows, widths, header_fill=fill_req, wrap_cols=None):
    for j, h in enumerate(headers, 1):
        c = ws.cell(row=start_row, column=j, value=h)
        c.font = f_head
        c.fill = header_fill[j - 1] if isinstance(header_fill, list) else header_fill
        c.alignment = center_wrap
        c.border = border
    for i, r in enumerate(rows, start_row + 1):
        for j, v in enumerate(r, 1):
            c = ws.cell(row=i, column=j, value=v)
            c.font = f_norm
            c.alignment = wrap
            c.border = border
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)


rows = D.build_rows()
wb = Workbook()

# ------------------------------------------------------------------ Tu_vung (tạo trước để biết vị trí)
vocab_keys = ["speech_act", "urgency", "language_mix", "style", "speaking_rate", "voice_persona", "driving_state",
              "cabin_type", "windows", "music", "rain", "playback_channel", "status", "norm_features", "eval_focus",
              "data_origin", "voice_source", "noise_source", "rir_source", "lombard_applied", "qc_check", "qc_status",
              "trigger_type", "lexicon_variant", "gender", "age_group", "region_dialect", "dialect_intensity", "voice_quality", "f0_band"]
VCOL = {k: get_column_letter(i + 1) for i, k in enumerate(vocab_keys)}
VLEN = {k: len(D.VOCAB[k]) for k in vocab_keys}

# ------------------------------------------------------------------ Tong_quan
ws = wb.active
ws.title = "Tong_quan"
ws.sheet_view.showGridLines = False
ws.column_dimensions["A"].width = 26
ws.column_dimensions["B"].width = 78
ws.column_dimensions["C"].width = 16
ws.column_dimensions["D"].width = 50


def sec(r, text):
    c = ws.cell(row=r, column=1, value=text)
    c.font = f_bold
    for j in range(1, 5):
        ws.cell(row=r, column=j).fill = fill_sec


ws["A1"] = "Bộ dữ liệu mẫu TTS cho trợ lý ảo trên ô tô (tiếng Việt)"
ws["A1"].font = f_title
ws["A2"] = "Phiên bản mẫu 0.2 – 11/10/2026. Dùng để thống nhất cấu trúc trước khi tạo bộ dữ liệu thật bằng AI tạo sinh, vì nhiều yếu tố giọng nói và môi trường xe không thu trực tiếp được."
ws["A2"].font = f_norm

sec(4, "MỤC ĐÍCH")
ws["A5"] = "Nội dung"
NS = len(D.build_rows())
FULL = TX.full_table()
NF = len(D.FIELDS)
ws["B5"] = (f"Cung cấp lược đồ, 12 case study, 17 preset môi trường, bảng phân loại {len(FULL)} yếu tố ảnh hưởng thành 6 nhóm "
            f"(nhóm Người nói gồm giọng vùng miền, từ ngữ địa phương, tần số giọng; 11 yếu tố môi trường mới ở phiên bản 0.2), "
            f"giải pháp cho từng nhóm khi không thu được dữ liệu thật, quy trình tạo dữ liệu tổng hợp, hồ sơ người nói, bảng từ vựng địa phương và {NS} câu mẫu.")
sec(7, "CÁCH DÙNG")
steps = [
    ("Bước 1", f"Xem sheet Luoc_do để biết ý nghĩa {NF} trường; Case_study và Preset_moi_truong để chọn tình huống và điều kiện; Phan_loai_yeu_to để biết các nhóm yếu tố và khả năng lấy dữ liệu thật; Giai_phap để biết cách xử lý khi không thu được."),
    ("Bước 2", "Thêm câu mới vào cuối sheet Mau_du_lieu; các cột có mũi tên chọn sẽ kiểm tra giá trị hợp lệ theo sheet Tu_vung."),
    ("Bước 3", "Các dòng mẫu là câu ví dụ; thay hoặc xóa khi nhập dữ liệu thật và đổi trạng thái sang “Chờ thu”. Chọn speaker_id từ sheet Ho_so_nguoi_noi; ba cột speaker_gender, speaker_dialect, speaker_f0_band tự tra theo mã đó."),
    ("Bước 4", "Cột est_duration_s, asr_input và ba cột speaker_* tự tính; đừng nhập tay. Các cột xanh lá (nguồn gốc, kiểm tra chất lượng) do pipeline điền."),
    ("Bước 5", f"File .jsonl cùng thư mục chứa đúng {NS} dòng mẫu, cùng tên trường, để đưa thẳng vào pipeline."),
]
for i, (a, b) in enumerate(steps, 8):
    ws.cell(row=i, column=1, value=a).font = f_bold
    c = ws.cell(row=i, column=2, value=b)
    c.font = f_norm
    c.alignment = wrap

sec(14, "CHÚ THÍCH MÀU")
ws["A15"] = "Tiêu đề xanh đậm"; ws["A15"].fill = fill_req; ws["A15"].font = f_head
ws["B15"] = "Cột bắt buộc điền."; ws["B15"].font = f_norm
ws["A16"] = "Tiêu đề xám"; ws["A16"].fill = fill_opt; ws["A16"].font = f_head
ws["B16"] = "Cột tùy chọn."; ws["B16"].font = f_norm
ws["A17"] = "Chữ xanh dương, nền vàng"; ws["A17"].fill = fill_input; ws["A17"].font = f_input
ws["B17"] = "Ô tham số có thể sửa (ô C20:C22 bên dưới)."; ws["B17"].font = f_norm
ws["A18"] = "Tiêu đề xanh lá"; ws["A18"].fill = fill_prov; ws["A18"].font = f_head
ws["B18"] = "Cột nguồn gốc và kiểm tra chất lượng của dữ liệu tổng hợp; pipeline điền, người dùng không nhập tay."; ws["B18"].font = f_norm

sec(19, "THAM SỐ ƯỚC LƯỢNG THỜI LƯỢNG (có thể sửa)")
for r, (name, val) in zip((20, 21, 22), (("Chậm", 3.6), ("Vừa", 4.5), ("Nhanh", 5.4))):
    ws.cell(row=r, column=1, value=name).font = f_bold
    ws.cell(row=r, column=2, value="Số âm tiết mỗi giây (giả định tự đặt, chưa đo; hiệu chỉnh theo giọng thật)").font = f_norm
    c = ws.cell(row=r, column=3, value=val)
    c.font = f_input
    c.fill = fill_input
    c.number_format = "0.0"
RATE_CELL = {"Chậm": "Tong_quan!$C$20", "Vừa": "Tong_quan!$C$21", "Nhanh": "Tong_quan!$C$22"}

sec(24, "SỐ MẪU THEO CASE STUDY (tự đếm)")
for j, h in enumerate(["case_id", "case_name", "Số mẫu"], 1):
    c = ws.cell(row=25, column=j, value=h)
    c.font = f_head; c.fill = fill_req; c.border = border; c.alignment = center_wrap
for k in range(12):
    r = 26 + k
    ws.cell(row=r, column=1, value=f"=Case_study!A{k + 2}").font = f_norm
    ws.cell(row=r, column=2, value=f"=Case_study!B{k + 2}").font = f_norm
    ws.cell(row=r, column=3, value=f"=COUNTIF(Mau_du_lieu!$B$2:$B$1000,A{r})").font = f_norm
    for j in range(1, 4):
        ws.cell(row=r, column=j).border = border
ws.cell(row=38, column=1, value="Tổng").font = f_bold
ws.cell(row=38, column=3, value="=SUM(C26:C37)").font = f_bold

sec(40, "LƯU Ý VỀ GIÁ TRỊ MÔI TRƯỜNG")
ws["A41"] = "Giả định"
ws["B41"] = ("Các khoảng ồn nền dB(A) và SNR mục tiêu là giá trị chỉ dẫn từ hiểu biết chung, chưa đo trên xe cụ thể. "
             "Cần đo thực tế (mic đặt tại tai người lái, từng mẫu xe, từng tốc độ) rồi thay vào Preset_moi_truong trước khi dùng để thu hoặc mô phỏng.")
ws["B41"].alignment = wrap; ws["B41"].font = f_norm
for r in (5, 41):
    ws.row_dimensions[r].height = 58
for r in range(8, 13):
    ws.row_dimensions[r].height = 42

sec(43, "SỐ YẾU TỐ THEO NHÓM (tự đếm từ sheet Phan_loai_yeu_to)")
for j, h in enumerate(["Nhóm", "", "Số yếu tố"], 1):
    c = ws.cell(row=44, column=j, value=h)
    c.font = f_head; c.fill = fill_req; c.border = border; c.alignment = center_wrap
for k, g in enumerate([TX.G1, TX.G2, TX.G3, TX.G4, TX.G5, TX.G6]):
    r = 45 + k
    ws.cell(row=r, column=1, value=g).font = f_norm
    ws.cell(row=r, column=3, value=f"=COUNTIF(Phan_loai_yeu_to!$A$2:$A$200,A{r})").font = f_norm
    for j in range(1, 4):
        ws.cell(row=r, column=j).border = border
ws.cell(row=51, column=1, value="Tổng").font = f_bold
ws.cell(row=51, column=3, value="=SUM(C45:C50)").font = f_bold

sec(53, "SƠ ĐỒ CÁC SHEET")
SHEETMAP = [
    ("Luoc_do", "Ý nghĩa từng trường của Mau_du_lieu"),
    ("Case_study", "12 tình huống; có cột tự tính mức liên quan nhận dạng giọng nói (ASR)"),
    ("Mau_du_lieu", "Câu mẫu với đầy đủ nhãn môi trường, người nói, nguồn gốc và kiểm tra chất lượng"),
    ("Preset_moi_truong", "Các điều kiện môi trường dựng sẵn"),
    ("Phan_loai_yeu_to", "Bảng yếu tố đầy đủ theo 6 nhóm, kèm khả năng lấy dữ liệu thật và nhóm giải pháp"),
    ("Giai_phap", "Cách tạo dữ liệu khi không thu thật, công cụ, cổng kiểm tra, việc thu thật tối thiểu và rủi ro"),
    ("Quy_trinh_tong_hop", "Chín bước tạo dữ liệu tổng hợp và cổng kiểm tra từng bước"),
    ("Nguon_giay_phep", "Nguồn dữ liệu, mô hình, công cụ và tình trạng kiểm chứng giấy phép"),
    ("Ho_so_nguoi_noi", "Hồ sơ giọng: giới, tuổi, vùng, tần số giọng, chất giọng"),
    ("Tu_dia_phuong", "Từ vựng theo vùng, cần người bản ngữ duyệt"),
    ("Tu_vung", "Danh sách giá trị hợp lệ cho các ô chọn"),
]
for k, (a, b) in enumerate(SHEETMAP):
    ws.cell(row=54 + k, column=1, value=a).font = f_bold
    c = ws.cell(row=54 + k, column=2, value=b); c.font = f_norm; c.alignment = wrap

sec(66, "NGUYÊN TẮC VỚI DỮ LIỆU TỔNG HỢP")
PRINC = [
    "Giữ cột data_origin cho mọi dòng và luôn giữ một tập thu thật riêng, không dùng để huấn luyện, làm mốc đo khoảng cách giữa tổng hợp và thực tế.",
    "Giọng tạo từ người thật chỉ dùng khi có đồng ý bằng văn bản; kiểm giấy phép mô hình sinh giọng trước khi dùng dữ liệu cho sản phẩm thương mại (xem Nguon_giay_phep).",
    "Dùng nhiều mô hình sinh giọng khác nhau và nhiều bộ ồn khác nhau để model không học dấu vết của một bộ sinh.",
    "Nếu mô hình cuối là TTS, dữ liệu tổng hợp từ TTS khác bị giới hạn bởi chất lượng mô hình đó; dùng chủ yếu cho kịch bản đánh giá, ASR và phủ điều kiện, và so với tập thu thật.",
]
for k, t in enumerate(PRINC):
    ws.cell(row=67 + k, column=1, value=f"Nguyên tắc {k + 1}").font = f_bold
    c = ws.cell(row=67 + k, column=2, value=t); c.font = f_norm; c.alignment = wrap
    ws.row_dimensions[67 + k].height = 44

# ------------------------------------------------------------------ Luoc_do
ws2 = wb.create_sheet("Luoc_do")
write_table(ws2, 1, ["Tên trường (snake_case)", "Nhãn tiếng Việt", "Kiểu", "Bắt buộc", "Mô tả", "Giá trị / ví dụ"],
            [(f[0], f[1], f[2], "Có" if f[3] else "Không", f[4], f[5]) for f in D.FIELDS],
            [24, 30, 20, 10, 60, 44])

# ------------------------------------------------------------------ Case_study
ws3 = wb.create_sheet("Case_study")
write_table(ws3, 1, ["case_id", "case_name", "Mô tả tình huống", "Ý định tiêu biểu", "Yêu cầu TTS trọng yếu",
                     "Rủi ro / lỗi cần tránh", "Preset môi trường ưu tiên", "Ưu tiên dữ liệu"],
            [list(c) for c in D.CASES], [10, 28, 44, 44, 56, 52, 22, 14])

# ------------------------------------------------------------------ Mau_du_lieu
ws4 = wb.create_sheet("Mau_du_lieu")
KEYS = D.FIELD_KEYS
col_letter = {k: get_column_letter(i + 1) for i, k in enumerate(KEYS)}
PROV = {"data_origin", "voice_source", "generator_model", "noise_source", "rir_source", "lombard_applied", "snr_actual_dB",
        "qc_asr_wer_pct", "qc_dialect_match", "qc_f0_in_band", "qc_native_review", "qc_status"}
header_fills = [fill_prov if k in PROV else (fill_req if k in D.REQUIRED else fill_opt) for k in KEYS]
widths = {"sample_id": 11, "case_id": 9, "case_name": 22, "intent": 26, "speech_act": 14, "urgency": 12,
          "text_raw": 52, "text_spoken": 60, "language_mix": 12, "norm_features": 30, "style": 18,
          "speaking_rate": 12, "voice_persona": 16, "est_duration_s": 12, "env_id": 9, "driving_state": 12,
          "speed_kmh": 10, "cabin_noise_dBA": 14, "snr_target_dB": 11, "cabin_type": 14, "hvac_level": 10,
          "windows": 14, "music": 9, "passengers": 11, "rain": 10, "playback_channel": 22, "eval_focus": 30,
          "status": 10, "notes": 44, "speaker_id": 11, "lexicon_variant": 14, "speaker_gender": 12,
          "speaker_dialect": 18, "speaker_f0_band": 12, "trigger_type": 22, "asr_input": 12, "data_origin": 20, "voice_source": 28,
          "generator_model": 20, "noise_source": 22, "rir_source": 22, "lombard_applied": 12, "snr_actual_dB": 12,
          "qc_asr_wer_pct": 12, "qc_dialect_match": 14, "qc_f0_in_band": 14, "qc_native_review": 14, "qc_status": 14}
DERIVED = {"speaker_gender": "C", "speaker_dialect": "E", "speaker_f0_band": "G"}
table_rows = []
for i, r in enumerate(rows, 2):
    row = []
    for k in KEYS:
        if k == "est_duration_s":
            T = f"{col_letter['text_spoken']}{i}"
            Rt = f"{col_letter['speaking_rate']}{i}"
            row.append(
                f'=IF({T}="","",ROUND((LEN(TRIM({T}))-LEN(SUBSTITUTE(TRIM({T})," ",""))+1)/'
                f'IF({Rt}="Chậm",{RATE_CELL["Chậm"]},IF({Rt}="Nhanh",{RATE_CELL["Nhanh"]},{RATE_CELL["Vừa"]})),1))')
        elif k == "asr_input":
            tc = f"{col_letter['trigger_type']}{i}"
            row.append(f'=IF({tc}="","",IF({tc}="Hệ thống chủ động","Không","Có"))')
        elif k in DERIVED:
            sid = f"{col_letter['speaker_id']}{i}"
            colx = DERIVED[k]
            row.append(f'=IF({sid}="","",IFERROR(INDEX(Ho_so_nguoi_noi!${colx}$2:${colx}$13,MATCH({sid},Ho_so_nguoi_noi!$A$2:$A$13,0)),""))')
        else:
            row.append(r[k])
        pass
    table_rows.append(row)
write_table(ws4, 1, KEYS, table_rows, [widths[k] for k in KEYS], header_fill=header_fills)
ws4.freeze_panes = "D2"
for i in range(2, len(rows) + 2):
    ws4.cell(row=i, column=KEYS.index("est_duration_s") + 1).number_format = "0.0"

# cột liên quan ASR trong Case_study (tự đếm từ Mau_du_lieu)
for j, h in ((9, "Số mẫu"), (10, "Số mẫu có đầu vào giọng nói"), (11, "Liên quan nhận dạng giọng nói (ASR)")):
    c = ws3.cell(row=1, column=j, value=h)
    c.font = f_head; c.fill = fill_opt; c.alignment = center_wrap; c.border = border
asr_col = col_letter["asr_input"]
for k in range(12):
    r = k + 2
    ws3.cell(row=r, column=9, value=f'=COUNTIF(Mau_du_lieu!$B$2:$B$1000,A{r})')
    ws3.cell(row=r, column=10, value=f'=COUNTIFS(Mau_du_lieu!$B$2:$B$1000,A{r},Mau_du_lieu!${asr_col}$2:${asr_col}$1000,"Có")')
    ws3.cell(row=r, column=11, value=f'=IF(I{r}=0,"",IF(J{r}=0,"Không",IF(J{r}=I{r},"Có","Một phần")))')
    for j in (9, 10, 11):
        c = ws3.cell(row=r, column=j); c.font = f_norm; c.border = border; c.alignment = wrap
ws3.column_dimensions["I"].width = 10; ws3.column_dimensions["J"].width = 16; ws3.column_dimensions["K"].width = 22
ws3.cell(row=15, column=1, value=("Ghi chú: cột I–K tự tính từ Mau_du_lieu. “Không” nghĩa là mọi mẫu của case này do hệ thống tự phát, "
          "không có lời người dùng; khi đó các yếu tố phía ASR (giọng, từ địa phương, trạng thái người nói) không áp dụng, chỉ còn yếu tố phía TTS.")).font = f_norm

# kiểm tra giá trị hợp lệ
dv_map = {k: k for k in ["speech_act", "urgency", "language_mix", "style", "speaking_rate", "voice_persona",
                         "driving_state", "cabin_type", "windows", "music", "rain", "playback_channel", "status", "lexicon_variant", "trigger_type",
                         "data_origin", "voice_source", "noise_source", "rir_source", "lombard_applied", "qc_status"]}
dv_map.update({"qc_dialect_match": "qc_check", "qc_f0_in_band": "qc_check", "qc_native_review": "qc_check"})
for k, vk in dv_map.items():
    rng = f"Tu_vung!${VCOL[vk]}$2:${VCOL[vk]}${VLEN[vk] + 1}"
    dv = DataValidation(type="list", formula1=f"={rng}", allow_blank=True, showErrorMessage=True,
                        errorTitle="Giá trị không hợp lệ", error="Chọn một giá trị trong danh sách (xem sheet Tu_vung).")
    ws4.add_data_validation(dv)
    dv.add(f"{col_letter[k]}2:{col_letter[k]}1000")
dv = DataValidation(type="list", formula1="=Ho_so_nguoi_noi!$A$2:$A$13", allow_blank=True, showErrorMessage=True,
                    errorTitle="Giá trị không hợp lệ", error="Chọn speaker_id trong sheet Ho_so_nguoi_noi.")
ws4.add_data_validation(dv); dv.add(f"{col_letter['speaker_id']}2:{col_letter['speaker_id']}1000")
dv = DataValidation(type="list", formula1="=Case_study!$A$2:$A$13", allow_blank=True, showErrorMessage=True,
                    errorTitle="Giá trị không hợp lệ", error="Chọn case_id trong sheet Case_study.")
ws4.add_data_validation(dv); dv.add(f"{col_letter['case_id']}2:{col_letter['case_id']}1000")
dv = DataValidation(type="list", formula1=f"=Preset_moi_truong!$A$2:$A${len(D.PRESETS) + 1}", allow_blank=True, showErrorMessage=True,
                    errorTitle="Giá trị không hợp lệ", error="Chọn env_id trong sheet Preset_moi_truong.")
ws4.add_data_validation(dv); dv.add(f"{col_letter['env_id']}2:{col_letter['env_id']}1000")

# ------------------------------------------------------------------ Preset_moi_truong
ws5 = wb.create_sheet("Preset_moi_truong")
write_table(ws5, 1, ["env_id", "Tên preset", "Mô tả", "Trạng thái lái", "Tốc độ (km/h)", "Ồn nền chỉ dẫn (dB(A))",
                     "Nguồn ồn chính", "SNR mục tiêu (dB)", "Mức quạt (0–5)", "Cửa kính", "Nhạc", "Mưa",
                     "Số người", "Mục tiêu đánh giá"],
            [list(p) for p in D.PRESETS], [9, 30, 40, 14, 12, 18, 38, 14, 12, 14, 9, 10, 10, 44])
ws5.cell(row=len(D.PRESETS) + 3, column=1, value="Ghi chú: ồn nền và SNR là giá trị chỉ dẫn/thiết kế, chưa đo; thay bằng số đo thực khi có.").font = f_norm

# ------------------------------------------------------------------ Phan_loai_yeu_to (gộp bảng yếu tố cũ)
ws8 = wb.create_sheet("Phan_loai_yeu_to")
write_table(ws8, 1, ["Nhóm", "Nhóm con", "Mã", "Yếu tố", "Mức / giá trị tiêu biểu", "Phía bị ảnh hưởng", "Tác động chính",
                     "Cách ghi nhận hoặc mô phỏng", "Chỉ số đánh giá", "Ưu tiên", "Trường dữ liệu tương ứng", "Nhóm giải pháp",
                     "Khả năng lấy dữ liệu thật", "Phiên bản"],
            [list(x) for x in FULL], [34, 36, 7, 34, 54, 16, 56, 46, 40, 10, 34, 12, 30, 12])
n_tax = len(FULL)
ws8.auto_filter.ref = f"A1:N{n_tax + 1}"
ws8.cell(row=n_tax + 3, column=1, value=("Ghi chú: “(đề xuất)” là trường chưa có trong Mau_du_lieu, chỉ nên thêm khi bộ dữ liệu cần. "
          "Phía bị ảnh hưởng: ASR là giọng người dùng đi vào hệ thống; TTS là giọng trợ lý đi ra loa. "
          "Yếu tố ghi “Đầu vào (ASR)” chỉ áp dụng cho mẫu có asr_input = Có. "
          "Mã E01 (giọng vùng miền) của bảng cũ được tách thành S41–S45. Cột “Nhóm giải pháp” trỏ sang sheet Giai_phap. "
          "Bảng Yeu_to_moi_truong của phiên bản 0.1 đã được gộp vào đây.")).font = f_norm

# ------------------------------------------------------------------ Giai_phap
ws11 = wb.create_sheet("Giai_phap")
sol_rows = []
for sol in TX.SOLUTIONS:
    codes = "; ".join(sorted(x[2] for x in FULL if x[11] == sol[0]))
    sol_rows.append([sol[0], sol[1], codes] + list(sol[2:]))
write_table(ws11, 1, ["Mã", "Nhóm giải pháp", "Yếu tố liên quan", "Hạn chế khi không thu được thật", "Giải pháp chính",
                      "Công cụ hoặc nguồn", "Cổng kiểm tra chất lượng", "Việc thu thật tối thiểu nên làm", "Rủi ro còn lại"],
            sol_rows, [7, 30, 34, 40, 80, 44, 52, 40, 52])
ws11.cell(row=len(sol_rows) + 3, column=1, value=("Ghi chú: các số lượng nêu ở cột “Việc thu thật tối thiểu” (ví dụ 10–20 tình nguyện viên) là đề xuất của tôi, "
          "chưa có nguồn nào đã đọc xác nhận con số này. Công cụ chưa kiểm tra giấy phép được ghi rõ ở sheet Nguon_giay_phep.")).font = f_norm

# ------------------------------------------------------------------ Quy_trinh_tong_hop
ws12 = wb.create_sheet("Quy_trinh_tong_hop")
write_table(ws12, 1, ["Bước", "Tên bước", "Đầu vào", "Việc làm", "Đầu ra", "Cổng kiểm tra", "Trường ghi lại"],
            [list(x) for x in TX.PIPELINE], [7, 30, 40, 80, 28, 56, 44])

# ------------------------------------------------------------------ Nguon_giay_phep
ws13 = wb.create_sheet("Nguon_giay_phep")
write_table(ws13, 1, ["Tên", "Loại", "Dùng cho", "Giấy phép hoặc điều kiện", "Mức kiểm chứng", "Ghi chú"],
            [list(x) for x in TX.SOURCES], [44, 30, 32, 44, 34, 52])
ws13.cell(row=len(TX.SOURCES) + 3, column=1, value=("Ghi chú: “Chưa kiểm tra” nghĩa là tôi chưa đọc giấy phép của mục đó; phải kiểm trước khi dùng cho sản phẩm thương mại. "
          "Các mục “Đã kiểm” lấy từ Bảng 1 của bài PhoASR, trang mô hình XTTS-v2 hoặc chính bài PhoAudiobook.")).font = f_norm

# ------------------------------------------------------------------ Ho_so_nguoi_noi
ws9 = wb.create_sheet("Ho_so_nguoi_noi")
write_table(ws9, 1, ["speaker_id", "Vai trò", "gender", "age_group", "region_dialect", "dialect_intensity", "f0_band",
                     "f0_range_hz (chỉ dẫn)", "f0_measured_hz (điền khi đo)", "voice_quality", "speaking_rate_default", "Ghi chú"],
            [[s[0], s[1], s[2], s[3], s[4], s[5], s[6], s[7], None, s[8], s[9], s[10]] for s in D.SPEAKERS],
            [11, 18, 9, 22, 28, 16, 12, 18, 18, 18, 18, 56],
            header_fill=[fill_req] * 8 + [fill_opt] + [fill_req] * 2 + [fill_opt])
for i in range(2, len(D.SPEAKERS) + 2):
    c = ws9.cell(row=i, column=9)
    c.fill = fill_input
    c.font = f_input
r0 = len(D.SPEAKERS) + 4
ws9.cell(row=r0, column=1, value="Ngưỡng dải F0 (chỉ dẫn, theo hiểu biết ngữ âm học chung; cần đo trên giọng thật)").font = f_bold
for j, h in enumerate(["f0_band", "Khoảng F0", "Nhóm tiêu biểu"], 1):
    c = ws9.cell(row=r0 + 1, column=j, value=h)
    c.font = f_head; c.fill = fill_req; c.border = border; c.alignment = center_wrap
for k, (a, b, c_) in enumerate(D.F0_BAND_GUIDE):
    for j, v in enumerate((a, b, c_), 1):
        cc = ws9.cell(row=r0 + 2 + k, column=j, value=v)
        cc.font = f_norm; cc.border = border
ws9.cell(row=r0 + 7, column=1, value=("Ô nền vàng (f0_measured_hz) để trống: điền giá trị F0 trung bình đo được. "
          "SPK001–SPK007 là giọng thu cho TTS; SPK101–SPK105 là ví dụ người dùng cho phía nhận dạng.")).font = f_norm
for fk, col in (("gender", "C"), ("age_group", "D"), ("region_dialect", "E"), ("dialect_intensity", "F"),
                ("f0_band", "G"), ("voice_quality", "J"), ("speaking_rate", "K")):
    rng = f"Tu_vung!${VCOL[fk]}$2:${VCOL[fk]}${VLEN[fk] + 1}"
    dv = DataValidation(type="list", formula1=f"={rng}", allow_blank=True, showErrorMessage=True,
                        errorTitle="Giá trị không hợp lệ", error="Chọn một giá trị trong danh sách (xem sheet Tu_vung).")
    ws9.add_data_validation(dv)
    dv.add(f"{col}2:{col}13")

# ------------------------------------------------------------------ Tu_dia_phuong
ws10 = wb.create_sheet("Tu_dia_phuong")
write_table(ws10, 1, ["Nhóm", "Khái niệm", "Chuẩn chung / miền Bắc", "Miền Trung", "Miền Nam", "Dùng cho", "Trạng thái duyệt"],
            [[a, b, c, d or "(chưa xác minh)", e, f, "Cần người bản ngữ duyệt"] for a, b, c, d, e, f in TX.LEXICON],
            [22, 28, 28, 22, 24, 36, 26])
ws10.cell(row=len(TX.LEXICON) + 3, column=1, value=("Ghi chú: đây là các ví dụ phổ biến dựa trên hiểu biết chung về phương ngữ, chưa qua khảo sát. "
          "Ô “(chưa xác minh)” nghĩa là chưa biết chắc; không điền đoán. Mỗi vùng cần ít nhất một người bản ngữ duyệt trước khi dùng thu âm.")).font = f_norm

# ------------------------------------------------------------------ Tu_vung
ws7 = wb.create_sheet("Tu_vung")
for j, k in enumerate(vocab_keys, 1):
    c = ws7.cell(row=1, column=j, value=k)
    c.font = f_head; c.fill = fill_req; c.alignment = center_wrap; c.border = border
    for i, v in enumerate(D.VOCAB[k], 2):
        cc = ws7.cell(row=i, column=j, value=v)
        cc.font = f_norm; cc.border = border; cc.alignment = wrap
    ws7.column_dimensions[get_column_letter(j)].width = 28
ws7.freeze_panes = "A2"

wb._sheets = [wb["Tong_quan"], wb["Luoc_do"], wb["Case_study"], wb["Mau_du_lieu"], wb["Preset_moi_truong"],
              wb["Phan_loai_yeu_to"], wb["Giai_phap"], wb["Quy_trinh_tong_hop"], wb["Nguon_giay_phep"],
              wb["Ho_so_nguoi_noi"], wb["Tu_dia_phuong"], wb["Tu_vung"]]
wb.save(OUT_XLSX)

# ------------------------------------------------------------------ JSONL
with open(OUT_JSONL, "w", encoding="utf-8") as fh:
    for r in rows:
        rec = {k: r[k] for k in KEYS}
        rec["norm_features"] = [x.strip() for x in r["norm_features"].split(";") if x.strip()]
        fh.write(json.dumps(rec, ensure_ascii=False) + "\n")
print("ok", len(rows))
