# -*- coding: utf-8 -*-
"""Tạo bảng nhãn nơi xử lý theo ý định (tại biên / cloud / cả hai) cho bộ v3.1 (đề xuất 28/10/2026, chờ duyệt).
Chỉ đọc data/m1m2_10k_v31 và gate_hop_v1.csv của lượt full10k; ghi data/nhan_xu_ly_y_dinh/{.tsv,.xlsx}.
Chạy (từ gen10k/): py -3.13 -m phantich.nhan_xu_ly_y_dinh
Sau khi người dùng sửa xlsx, xlsx là bản chuẩn; sửa bảng N ở đây cho khớp trước khi chạy lại (chạy lại sẽ ghi đè xlsx)."""
import collections
import csv
import glob
import json
import os

from cauhinh.duong_dan import ROOT as G
B, C, H = "bien", "cloud", "ca_hai"
TEN = {B: "Tại biên", C: "Cloud", H: "Cả hai (có mạng tốt hơn)"}
OK, DUYET = "khá chắc", "cần duyệt"

# (case, intent) -> (nhãn, khi mất mạng, lý do, mức chắc)
N = {
    # CS01 dẫn đường: giả định xe có bản đồ ngoại tuyến (embedded); giao thông thời gian thực cần mạng
    ("CS01", "chon_app_ban_do"): (H, "Mở/đổi app trên xe được; app bản đồ online (Google Maps, Waze) không tìm được đường", "Đổi app là lệnh trên xe; nội dung app phần lớn cần mạng", DUYET),
    ("CS01", "chon_lo_trinh"): (H, "Tính lộ trình bằng bản đồ ngoại tuyến; không có thông tin tắc đường", "Ràng buộc lộ trình (tránh cao tốc, đi đường X) xử lý được trên xe", OK),
    ("CS01", "dan_duong_den_dia_chi"): (H, "Tìm địa chỉ trong bản đồ ngoại tuyến, có thể thiếu số nhà/hẻm mới", "Mã hóa địa chỉ trực tuyến đầy đủ hơn", OK),
    ("CS01", "dan_duong_den_dia_diem"): (H, "Chỉ tìm được địa điểm có trong CSDL POI trên xe", "POI trực tuyến mới và đầy đủ hơn", OK),
    ("CS01", "dan_duong_den_khu_vuc"): (H, "Quận/tỉnh có sẵn trong bản đồ ngoại tuyến", "Đơn vị hành chính ít đổi; tên sau sáp nhập 2025 cần cập nhật bản đồ", OK),
    ("CS01", "dieu_khien_ban_do"): (B, "Hoạt động bình thường", "Thao tác giao diện bản đồ trên xe", OK),
    ("CS01", "hoi_chi_dan_tiep_theo"): (B, "Hoạt động bình thường khi đang dẫn đường", "Lấy từ lộ trình đang chạy", OK),
    ("CS01", "hoi_con_bao_xa"): (H, "Tính theo bản đồ ngoại tuyến", "Khoảng cách tới điểm đến hiện tại tính trên xe; tới địa điểm bất kỳ cần tìm POI", OK),
    ("CS01", "hoi_thoi_gian_den"): (H, "Ước tính theo tốc độ trung bình, không tính tắc đường", "Giờ đến chính xác cần giao thông thời gian thực", OK),
    ("CS01", "huy_dan_duong"): (B, "Hoạt động bình thường", "Lệnh điều khiển dẫn đường", OK),
    ("CS01", "nhac_lai_chi_dan"): (B, "Hoạt động bình thường", "Lặp lại chỉ dẫn đã có", OK),
    ("CS01", "tat_bat_tieng_dan_duong"): (B, "Hoạt động bình thường", "Cài đặt âm thanh trên xe", OK),
    ("CS01", "them_diem_dung"): (H, "Thêm POI có trong CSDL ngoại tuyến", "Tìm điểm dừng dọc đường trực tuyến đầy đủ hơn", OK),
    ("CS01", "thu_phi_khong_dung"): (C, "Không xem được số dư, lịch sử trừ tiền", "Tài khoản VETC/ePass nằm trên máy chủ nhà cung cấp", OK),
    ("CS01", "tim_dia_diem_gan_day"): (H, "Chỉ tìm trong POI ngoại tuyến, không biết giờ mở cửa, còn chỗ", "Thông tin địa điểm thời gian thực cần mạng", OK),
    ("CS01", "tranh_ket_xe_cong_trinh"): (C, "Không biết chỗ tắc/công trình; chỉ đổi được lộ trình khác", "Cần dữ liệu giao thông thời gian thực", OK),
    ("CS01", "ve_nha_co_quan"): (H, "Điểm đã lưu + bản đồ ngoại tuyến", "Địa điểm đã lưu trên xe; tránh tắc cần mạng", OK),
    # CS02 hỗ trợ lái: cài đặt ADAS trên xe
    ("CS02", "bat_tat_he_thong_ho_tro"): (B, "Hoạt động bình thường", "Cài đặt ADAS, camera trên xe", OK),
    ("CS02", "canh_bao_diem_mu"): (B, "Hoạt động bình thường", "Cài đặt ADAS", OK),
    ("CS02", "canh_bao_lech_lan"): (B, "Hoạt động bình thường", "Cài đặt ADAS", OK),
    ("CS02", "canh_bao_met_moi"): (B, "Hoạt động bình thường (mở nhạc trực tuyến thì cần mạng)", "Phản hồi cảnh báo trên xe; vài câu kèm mở nhạc", OK),
    ("CS02", "canh_bao_va_cham"): (B, "Hoạt động bình thường", "Cài đặt ADAS", OK),
    ("CS02", "do_nhay_canh_bao_chung"): (B, "Hoạt động bình thường", "Cài đặt cảnh báo", OK),
    ("CS02", "ga_tu_dong"): (B, "Hoạt động bình thường", "Ga tự động, giới hạn tốc độ trên xe", OK),
    # CS03 điều khiển xe
    **{("CS03", i): (B, "Hoạt động bình thường", "Điều khiển thiết bị trên xe", OK) for i in (
        "che_do_lai", "cop_xe", "cua_kinh", "cua_so_troi", "den_noi_that", "dieu_hoa_bat_tat_nhiet_do", "dieu_hoa_che_do",
        "dieu_hoa_huong_gio", "dieu_hoa_lay_gio", "dieu_hoa_noi_nong_lanh", "dieu_hoa_quat_gio", "gat_mua", "ghe_nho_vi_tri",
        "ghe_suoi_lam_mat_massage", "guong", "khoa_mo_cua", "khoi_dong_xe", "say_kinh", "vo_lang_den_ngoai", "wifi_camera")},
    # CS04 giải trí
    ("CS04", "am_luong"): (B, "Hoạt động bình thường", "Âm lượng trên xe", OK),
    ("CS04", "chon_app_ban_do"): (B, "Hoạt động bình thường", "Cài app bản đồ mặc định, hỏi app đang dùng", OK),
    ("CS04", "chuyen_bai_tam_dung"): (B, "Điều khiển được; nguồn phát trực tuyến có thể dừng vì hết đệm", "Lệnh điều khiển trình phát", OK),
    ("CS04", "chuyen_man_hinh_app"): (B, "Hoạt động bình thường", "Giao diện trên xe", OK),
    ("CS04", "dieu_khien_video"): (H, "Tua, dừng, quay lại được; đề xuất kênh, video mới thì không", "Điều khiển phát trên xe; thao tác với tài khoản app cần mạng", DUYET),
    ("CS04", "hoi_bai_dang_phat"): (H, "Đọc được tên bài/ca sĩ từ thông tin bài; năm phát hành, bài khác thì không", "Thông tin mở rộng cần tra cứu trực tuyến", OK),
    ("CS04", "mo_app"): (B, "Mở được app; app trực tuyến có thể không có nội dung", "Mở/thoát app trên xe", OK),
    ("CS04", "ngu_canh_xe_do"): (H, "Mở được nội dung đã tải/USB; phim trực tuyến thì không", "Chủ yếu mở phim/video", DUYET),
    ("CS04", "nguon_nhac_bluetooth_usb"): (B, "Hoạt động bình thường", "Chọn nguồn phát trên xe", OK),
    ("CS04", "phat_bai_ca_si"): (H, "Chỉ phát được bài có trong USB/điện thoại/đã tải", "Phát theo tên bài trực tuyến cần mạng", OK),
    ("CS04", "phat_the_loai_playlist"): (H, "Chỉ phát danh sách đã tải hoặc thư viện trên máy", "Danh sách trực tuyến (YouTube, Spotify) cần mạng", OK),
    ("CS04", "phat_theo_tam_trang"): (H, "Chọn trong thư viện trên máy nếu có gắn thể loại", "Gợi ý theo tâm trạng thường do dịch vụ trực tuyến", DUYET),
    ("CS04", "podcast"): (C, "Chỉ phát được tập đã tải", "Podcast, sách nói chủ yếu phát trực tuyến", DUYET),
    ("CS04", "radio"): (B, "Hoạt động bình thường với đài FM/AM", "Bộ thu FM trên xe (đài Internet thì cần mạng)", OK),
    ("CS04", "tim_noi_dung_app"): (C, "Không tìm được", "Tìm kiếm trong YouTube, app nhạc là dịch vụ trực tuyến", OK),
    ("CS04", "yeu_thich"): (H, "Lưu vào danh sách trên máy; đồng bộ tài khoản khi có mạng", "Danh sách của app trực tuyến lưu trên máy chủ", DUYET),
    # CS05 điện thoại, tin nhắn: gọi điện/SMS qua điện thoại kết nối = sóng di động, không cần Internet
    ("CS05", "chia_se_vi_tri"): (H, "Gửi được bằng SMS (vị trí từ GPS); qua Zalo/Messenger thì không", "Vị trí tính trên xe; gửi qua app cần Internet", OK),
    ("CS05", "doc_thong_bao"): (B, "Đọc được thông báo đã nhận", "Thông báo đã có trên xe/điện thoại", OK),
    ("CS05", "doc_tin_nhan"): (B, "Đọc được SMS đã nhận", "Tin nhắn trên điện thoại kết nối", OK),
    ("CS05", "goi_lai"): (B, "Gọi qua sóng di động", "Danh bạ, nhật ký gọi trên điện thoại kết nối", OK),
    ("CS05", "goi_theo_so"): (B, "Gọi qua sóng di động", "Gọi điện thoại thường", OK),
    ("CS05", "goi_theo_ten"): (B, "Gọi qua sóng di động", "Tra danh bạ trên điện thoại kết nối", OK),
    ("CS05", "nhan_tu_choi_cuoc_goi"): (B, "Hoạt động bình thường", "Điều khiển cuộc gọi", OK),
    ("CS05", "nhan_zalo"): (C, "Không gửi/gọi được", "Zalo, Telegram, Messenger là app trực tuyến", OK),
    ("CS05", "soan_gui_tin_nhan"): (H, "Gửi SMS được; nội dung tự do do ASR trên xe nhận dạng, có thể kém hơn", "Đọc chính tả tự do thường dùng ASR đám mây để chính xác hơn", DUYET),
    ("CS05", "tin_nhan_app"): (C, "Chỉ đọc được tin đã nhận; trả lời thì không", "Tin nhắn app bên thứ ba (Zalo, Viber) cần Internet", OK),
    # CS06 trạng thái xe
    ("CS06", "ap_suat_lop"): (B, "Hoạt động bình thường", "Cảm biến áp suất lốp", OK),
    ("CS06", "bao_duong"): (B, "Hoạt động bình thường (đặt lịch xưởng thì cần mạng)", "Lịch bảo dưỡng, mức dung dịch, lời nhắc trên xe", OK),
    ("CS06", "den_canh_bao"): (H, "Giải thích theo sổ tay lưu trên xe", "Giải thích chi tiết, tra lỗi theo mã thì tốt hơn khi có mạng", DUYET),
    ("CS06", "nhien_lieu_quang_duong"): (B, "Hoạt động bình thường", "Dữ liệu xe", OK),
    ("CS06", "nhiet_do_dong_co"): (B, "Hoạt động bình thường", "Cảm biến xe", OK),
    ("CS06", "so_km_da_di"): (B, "Hoạt động bình thường", "Đồng hồ công-tơ-mét", OK),
    # CS07 sạc, nhiên liệu
    ("CS07", "cong_sac"): (B, "Hoạt động bình thường", "Điều khiển nắp/cổng sạc", OK),
    ("CS07", "gioi_han_sac"): (B, "Hoạt động bình thường", "Cài đặt sạc trên xe", OK),
    ("CS07", "hen_gio_sac"): (B, "Hoạt động bình thường", "Hẹn giờ sạc trên xe", OK),
    ("CS07", "phan_tram_pin"): (B, "Hoạt động bình thường; mức pin khi đến nơi tính theo bản đồ ngoại tuyến", "Dữ liệu pin trên xe", OK),
    ("CS07", "quang_duong_con_lai_ev"): (H, "Ước tính theo mức tiêu hao trung bình", "Ước tính chính xác hơn khi có giao thông, độ cao, thời tiết", OK),
    ("CS07", "thoi_gian_chi_phi_sac"): (H, "Thời gian sạc tính được; chi phí, giá điện của trạm thì không", "Giá sạc theo trạm cần dữ liệu trực tuyến", OK),
    ("CS07", "tim_tram_sac"): (H, "Tìm trạm trong CSDL ngoại tuyến; không biết trạm còn trống/hoạt động", "Tình trạng trạm sạc thời gian thực cần mạng", OK),
    ("CS07", "tim_tram_xang"): (H, "Tìm cây xăng ngoại tuyến; giá xăng thì không", "Giá xăng cần tin trực tuyến", OK),
    # CS08 giao thông, thời tiết
    ("CS08", "camera_phat_nguoi"): (H, "Báo vị trí camera có trong bản đồ ngoại tuyến", "Danh sách camera cập nhật trực tuyến", DUYET),
    ("CS08", "chat_luong_khong_khi"): (C, "Không trả lời được (trừ cảm biến PM2.5 trong xe nếu có)", "AQI là dữ liệu trực tuyến", OK),
    ("CS08", "gio_den_theo_giao_thong"): (C, "Chỉ ước tính không có tắc đường", "Câu hỏi nêu rõ theo giao thông", OK),
    ("CS08", "ngap_nuoc"): (C, "Không biết chỗ ngập", "Dữ liệu ngập thời gian thực", OK),
    ("CS08", "nhiet_do_mua"): (C, "Chỉ đọc nhiệt độ ngoài xe từ cảm biến; dự báo thì không", "Phần lớn là hỏi dự báo theo nơi và giờ", DUYET),
    ("CS08", "tac_duong_phia_truoc"): (C, "Không biết tình hình giao thông", "Giao thông thời gian thực", OK),
    ("CS08", "tai_nan_su_co"): (C, "Không biết sự cố", "Tin sự cố thời gian thực", OK),
    ("CS08", "thoi_tiet_hom_nay_mai"): (C, "Không trả lời được", "Dự báo thời tiết trực tuyến", OK),
    ("CS08", "toc_do_toi_da"): (B, "Theo biển báo (camera) hoặc bản đồ ngoại tuyến", "Giới hạn tốc độ và cảnh báo vượt tốc xử lý trên xe", DUYET),
    # CS09 khẩn cấp: bắt buộc chạy khi mất Internet (gọi qua sóng di động / eCall)
    **{("CS09", i): (B, "Phải hoạt động: gọi/nhắn qua sóng di động hoặc eCall", "Khẩn cấp không được phụ thuộc Internet", OK) for i in (
        "bao_tai_nan", "bao_vi_tri_khan_cap", "cuu_ho_giao_thong", "goi_cap_cuu_115", "goi_cong_an_113", "goi_cuu_hoa_114", "keu_cuu", "toi_on")},
    # CS10 hội thoại chủ động
    ("CS10", "bao_thuc_hen_gio"): (B, "Hoạt động bình thường", "Báo thức trên xe/điện thoại", OK),
    ("CS10", "chao_cam_on"): (B, "Hoạt động bình thường", "Câu xã giao, trả lời mẫu", OK),
    ("CS10", "dat_lich"): (H, "Đặt/đọc lịch lưu trên máy; đồng bộ lịch tài khoản khi có mạng", "Lịch Google/Outlook lưu trên máy chủ", OK),
    ("CS10", "hoi_dap_chung"): (C, "Không trả lời được", "Hỏi đáp kiến thức chung cần tra cứu/LLM trực tuyến", OK),
    ("CS10", "hoi_gio_ngay"): (B, "Hoạt động bình thường (âm lịch, múi giờ tính trên máy)", "Đồng hồ và lịch trên xe", OK),
    ("CS10", "hoi_quan_an"): (C, "Chỉ tìm được quán có trong POI ngoại tuyến; đặt bàn thì không", "Tìm quán, đặt bàn là dịch vụ trực tuyến", OK),
    ("CS10", "hoi_quan_co_cho_do"): (C, "Không biết chỗ đỗ còn trống", "Thông tin bãi đỗ trực tuyến", DUYET),
    ("CS10", "ke_chuyen_cuoi"): (H, "Kể từ kho truyện lưu sẵn", "Truyện mới cần mạng", DUYET),
    ("CS10", "nhac_viec"): (B, "Hoạt động bình thường", "Lời nhắc lưu trên xe/điện thoại", OK),
    ("CS10", "tin_tuc"): (C, "Không có tin", "Tin tức trực tuyến", OK),
    ("CS10", "tinh_toan_doi_tien"): (H, "Tính toán được; đổi ngoại tệ theo tỷ giá thì không", "Tỷ giá cần dữ liệu trực tuyến", OK),
    ("CS10", "tro_chuyen_tam_su"): (H, "Trả lời mẫu đơn giản", "Trò chuyện tự do tốt hơn với LLM trực tuyến", DUYET),
    # CS11 hành khách
    ("CS11", "che_do_yen_tinh"): (B, "Hoạt động bình thường", "Cài đặt âm thanh trên xe", OK),
    ("CS11", "dieu_hoa_khoang_sau"): (B, "Hoạt động bình thường", "Điều hòa trên xe", OK),
    ("CS11", "khoa_tre_em"): (B, "Hoạt động bình thường", "Khóa cửa, kính trên xe", OK),
    ("CS11", "ngu_canh_xe_do"): (H, "Mở được app; nhạc trực tuyến (Zing MP3, YouTube Music) không phát", "Mở app nhạc trực tuyến", DUYET),
    ("CS11", "nhac_day_an_toan_ghe_sau"): (B, "Hoạt động bình thường", "Cảm biến dây an toàn", OK),
    ("CS11", "tat_loa_sau"): (B, "Hoạt động bình thường", "Âm thanh khoang sau", OK),
    ("CS11", "tre_noi"): (H, "Lệnh trên xe và câu hỏi lộ trình vẫn trả lời được", "Trẻ hỏi đủ loại; câu hỏi kiến thức cần mạng", DUYET),
    ("CS11", "tre_xem_nhac_phim_truyen"): (H, "Phát được nội dung đã tải/USB", "Phim, nhạc, sách nói phần lớn phát trực tuyến", DUYET),
    # CS12 làm rõ, trả lời ngắn: xử lý hội thoại trên xe; kết quả cuối phụ thuộc lệnh ở lượt trước
    ("CS12", "chon_ket_qua"): (B, "Hoạt động bình thường (chọn trong danh sách đã hiện)", "Quản lý hội thoại; danh sách có thể đến từ tìm kiếm trực tuyến ở lượt trước", OK),
    ("CS12", "ngoai_tuyen"): (B, "Phải hoạt động: báo mất mạng, đề xuất cách khác", "Ý định nói về chính việc mất mạng (tải bản đồ thì cần mạng)", OK),
    ("CS12", "noi_lai"): (B, "Hoạt động bình thường", "Lặp lại câu trả lời trước", OK),
    ("CS12", "sua_lenh"): (B, "Hoạt động bình thường", "Sửa giá trị của lệnh trước; thực hiện tùy lệnh gốc", OK),
    ("CS12", "tra_loi_gia_tri"): (B, "Hoạt động bình thường", "Điền giá trị cho câu hỏi của trợ lý", OK),
    ("CS12", "tu_choi_huy_dung"): (B, "Hoạt động bình thường", "Quản lý hội thoại", OK),
    ("CS12", "xac_nhan_dong_y"): (B, "Hoạt động bình thường", "Quản lý hội thoại", OK),
    ("CS12", "xem_them_lua_chon"): (B, "Lọc trong danh sách đã có", "So khoảng cách trong kết quả đã hiện", DUYET),
}

rows = [json.loads(l) for f in sorted(glob.glob(os.path.join(G, "data", "m1m2_10k_v31", "*.jsonl"))) for l in open(f, encoding="utf-8")]
by_sid = {r["sample_id"]: r for r in rows}
grp = collections.defaultdict(list)
for r in rows:
    grp[(r["case_id"], r["intent_id"])].append(r)
missing = set(grp) - set(N)
extra = set(N) - set(grp)
assert not missing and not extra, (missing, extra)
v1 = collections.Counter()
for f in glob.glob(os.path.join(G, "tts_gen", "runs", "ket_qua_full10k", "runs", "full10k_p*", "gate_hop_v1.csv")):
    for r in csv.DictReader(open(f, encoding="utf-8-sig")):
        if r["ket_qua"] == "dat":
            x = by_sid[r["sample_id"]]
            v1[(x["case_id"], x["intent_id"])] += 1
case_name = json.load(open(os.path.join(G, "data", "m1m2_10k_v31", "manifest.json"), encoding="utf-8"))["case_names"]

out = []
for k in sorted(grp):
    rs = grp[k]
    lab, off, why, conf = N[k]
    ex = " / ".join(list(dict.fromkeys(x["text_raw"] for x in rs[::max(1, len(rs) // 3)]))[:3])
    out.append(dict(case_id=k[0], case=case_name.get(k[0], ""), intent_id=k[1], intent=rs[0]["intent"], so_cau_v31=len(rs),
                    so_cau_v1_dat=v1[k], nhan=lab, nhan_doc=TEN[lab], khi_mat_mang=off, ly_do=why, muc_chac=conf, vi_du=ex))

OUT = os.path.join(G, "data", "nhan_xu_ly_y_dinh")
os.makedirs(OUT, exist_ok=True)
cols = list(out[0])
with open(os.path.join(OUT, "nhan_xu_ly_y_dinh.tsv"), "w", encoding="utf-8", newline="\n") as f:
    f.write("\t".join(cols) + "\n")
    for o in out:
        f.write("\t".join(str(o[c]) for c in cols) + "\n")

# xlsx để duyệt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
wb = Workbook()
ws0 = wb.active
ws0.title = "Hướng dẫn"
huong_dan = [
    ["Nhãn nơi xử lý theo ý định (bộ v3.1, 10.000 câu) — đề xuất ngày 28/10/2026, chờ người dùng duyệt"],
    [""],
    ["Nhãn nói về việc THỰC HIỆN yêu cầu có cần Internet không (không nói về nơi chạy ASR)."],
    ["bien  = Tại biên: xe tự hiểu và thực hiện, không cần Internet."],
    ["cloud = Cloud: không có Internet thì không làm được (hoặc chỉ báo lỗi)."],
    ["ca_hai = Cả hai: chạy được trên xe, có mạng thì kết quả tốt hơn (dữ liệu mới, thời gian thực, đầy đủ hơn)."],
    [""],
    ["Quy ước:"],
    ["• Gọi điện, SMS qua điện thoại kết nối = tại biên (cần sóng di động, không cần Internet). Zalo, Messenger, Telegram, Viber = cloud."],
    ["• Dẫn đường giả định xe có bản đồ ngoại tuyến; giao thông, tắc đường, ngập, giá cả = cần mạng."],
    ["• Khẩn cấp (CS09) và báo mất mạng (CS12 ngoai_tuyen) bắt buộc chạy được tại biên."],
    ["• Gán theo ý định, nên vài câu trong cùng ý định có thể khác (ví dụ CS08 nhiet_do_mua: 'ngoài trời bao nhiêu độ' đọc cảm biến xe)."],
    ["• Cột 'muc_chac' = 'cần duyệt' là chỗ nên xem kỹ. Sửa cột 'nhan' bằng danh sách chọn; ghi lý do vào 'ghi_chu_nguoi_duyet'."],
    [""],
    ["Tổng hợp (theo bảng hiện tại)"],
    ["Nhãn", "Số ý định", "Số câu v3.1", "Số câu v1 dat"],
]
for row in huong_dan:
    ws0.append(row)
for lab in (B, H, C):
    ws0.append([TEN[lab],
                f'=COUNTIF(Nhan!$G:$G,"{lab}")',
                f'=SUMIF(Nhan!$G:$G,"{lab}",Nhan!$E:$E)',
                f'=SUMIF(Nhan!$G:$G,"{lab}",Nhan!$F:$F)'])
ws0.column_dimensions["A"].width = 120
ws = wb.create_sheet("Nhan")
hdr = cols[:7] + ["khi_mat_mang", "ly_do", "muc_chac", "vi_du", "ghi_chu_nguoi_duyet"]
ws.append(hdr)
for o in out:
    ws.append([o[c] for c in cols[:7]] + [o["khi_mat_mang"], o["ly_do"], o["muc_chac"], o["vi_du"], ""])
dv = DataValidation(type="list", formula1='"bien,ca_hai,cloud"', allow_blank=False)
ws.add_data_validation(dv)
dv.add(f"G2:G{len(out) + 1}")
fill = {B: "E2F0D9", H: "FFF2CC", C: "DDEBF7"}
for i, o in enumerate(out, start=2):
    ws.cell(i, 7).fill = PatternFill("solid", fgColor=fill[o["nhan"]])
    if o["muc_chac"] == DUYET:
        ws.cell(i, 10).font = Font(bold=True, color="C00000")
for c, w in zip("ABCDEFGHIJKL", (7, 24, 26, 30, 9, 9, 9, 50, 50, 11, 70, 30)):
    ws.column_dimensions[c].width = w
for row in ws.iter_rows(min_row=2):
    for cell in row:
        cell.alignment = Alignment(wrap_text=True, vertical="top")
for cell in ws[1]:
    cell.font = Font(bold=True)
ws.freeze_panes = "E2"
ws.auto_filter.ref = ws.dimensions
wb.save(os.path.join(OUT, "nhan_xu_ly_y_dinh.xlsx"))

t = collections.Counter()
for o in out:
    t[(o["nhan"], "y")] += 1; t[(o["nhan"], "c")] += o["so_cau_v31"]; t[(o["nhan"], "v")] += o["so_cau_v1_dat"]
for lab in (B, H, C):
    print(TEN[lab], t[(lab, "y")], "ý định,", t[(lab, "c")], "câu v3.1,", t[(lab, "v")], "câu v1 dat")
print("cần duyệt:", sum(o["muc_chac"] == DUYET for o in out), "ý định")
