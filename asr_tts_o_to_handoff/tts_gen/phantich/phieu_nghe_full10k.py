# -*- coding: utf-8 -*-
"""Phiếu nghe cho lượt 10.000 câu v3.1 (runs/ket_qua_full10k): câu GHÉP lời gọi + câu NGẪU NHIÊN.

Chạy (từ tts_gen/, venv có openpyxl):  .venv_zip/Scripts/python -m phantich.phieu_nghe_full10k
Ghi runs/nghe_thu_full10k.xlsx (+ .jsonl danh sách). Tệp phải nằm trong runs/ (liên kết tương đối tới ket_qua_full10k/).
Chọn câu (tất định theo băm, KHÔNG nhìn kết quả chấm của máy):
  G · Ghép lời gọi, 20 câu: 7 đủ cụm, 7 đuôi cắt đúng ranh giới âm tiết, 6 đuôi cắt giữa âm tiết (2 mỗi vị trí);
      chia đều 2 cụm "Trợ lý ơi"/"Hây trợ lý" và 3 vùng.
  N · Ngẫu nhiên, 100 câu từ 10.000 câu thường: chia đều 3 vùng, trải đều 10 phần.
Đợt 1 = 20 câu G + 30 câu N (khoảng 15–20 phút). Đợt 2, 3 = 35 câu N mỗi đợt (tùy chọn, để ước lượng chắc hơn).
Các cột máy (kết quả chấm v1/v2, câu L/M/Z/T nghe ra, giọng, mã) bị gập để người nghe không bị dẫn dắt.
"""
import collections
import csv
import hashlib
import json
import os
import time

from cauhinh.duong_dan import GOC as HERE  # noqa: E402
from thuvien.manifest import active_rows, load_manifest  # noqa: E402

ROOT = os.path.join(HERE, "runs", "ket_qua_full10k", "runs")
REL = "ket_qua_full10k/runs"
N_RANDOM = 100
WAKE_PLAN = [("du", "", 7), ("duoi", "ranh_gioi", 7), ("duoi", "giua_am_tiet_0", 2), ("duoi", "giua_am_tiet_1", 2),
             ("duoi", "giua_am_tiet_2", 2)]
VERDICT_VI = {"dat": "Đạt", "dat_1nguon": "Đạt một nguồn", "can_nghe": "Cần nghe", "loai": "Loại", "sinh_lai": "Sinh lại"}


def h(*p):
    return int(hashlib.sha1("|".join(map(str, p)).encode()).hexdigest()[:12], 16)


def gate(run, name):
    p = os.path.join(ROOT, run, f"{name}.csv")
    with open(p, encoding="utf-8-sig") as f:
        return {r["utt_id"]: r for r in csv.DictReader(f)}


def pick():
    out = []
    # G: câu ghép
    wm = active_rows(load_manifest(os.path.join(ROOT, "full10k_wake")))
    g1, g2 = gate("full10k_wake", "gate_hop_v1"), gate("full10k_wake", "gate_hop")
    for kind, cat, n in WAKE_PLAN:
        pool = [m for m in wm if m["wake_type"] == kind and (m["wake_cat"] or "") == cat]
        # xen kẽ 2 cụm và 3 vùng: (cụm, vùng) lần lượt theo vòng, trong mỗi ô lấy theo băm
        cells = collections.defaultdict(list)
        for m in sorted(pool, key=lambda m: h("G", m["utt_id"])):
            cells[(m["wake_phrase"], m["speaker_region"])].append(m)
        order = [(p, r) for r in ("Bắc", "Trung", "Nam") for p in ("tro_ly_oi", "hay_tro_ly")]
        order = order[h("lech", kind, cat) % len(order):] + order[:h("lech", kind, cat) % len(order)]
        got = []
        while len(got) < n and any(cells[k] for k in order):
            for k in order:
                if len(got) < n and cells[k]:
                    got.append(cells[k].pop(0))
        for m in got:
            out.append(dict(nhom="G", run="full10k_wake", m=m, g1=g1[m["utt_id"]], g2=g2[m["utt_id"]]))
    # N: ngẫu nhiên, chia đều vùng, trải đều phần
    per_reg = {"Bắc": N_RANDOM // 3 + 1, "Trung": N_RANDOM // 3, "Nam": N_RANDOM // 3}
    cand = collections.defaultdict(list)
    gates = {}
    for i in range(1, 11):
        run = f"full10k_p{i:02d}"
        gates[run] = (gate(run, "gate_hop_v1"), gate(run, "gate_hop"))
        for m in active_rows(load_manifest(os.path.join(ROOT, run))):
            cand[m["speaker_region"]].append((run, m))
    for reg, n in per_reg.items():
        rows = sorted(cand[reg], key=lambda x: h("N", x[1]["sample_id"]))
        by_run = collections.defaultdict(list)
        for run, m in rows:
            by_run[run].append(m)
        k = 0
        while k < n:                         # vòng tròn qua 10 phần
            for run in sorted(by_run):
                if k < n and by_run[run]:
                    m = by_run[run].pop(0)
                    out.append(dict(nhom="N", run=run, m=m, g1=gates[run][0][m["utt_id"]], g2=gates[run][1][m["utt_id"]]))
                    k += 1
    # đợt: 1 = G + 30 N; 2, 3 = phần N còn lại
    gs = [r for r in out if r["nhom"] == "G"]
    ns = sorted((r for r in out if r["nhom"] == "N"), key=lambda r: h("dot", r["m"]["sample_id"]))
    for i, r in enumerate(ns):
        r["dot"] = 1 if i < 30 else (2 if i < 65 else 3)
    for r in gs:
        r["dot"] = 1
    res = gs + ns
    res.sort(key=lambda r: (r["dot"], h("tron", r["m"]["utt_id"])))
    return res


def write(path, rows):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = Workbook()
    hd = wb.active
    hd.title = "Hướng dẫn"
    n_g = sum(r["nhom"] == "G" for r in rows)
    lines = [
        ("Nghe thử lượt 10.000 câu (v3.1, A100, 25/10/2026)", True),
        (f"{len(rows)} câu: {n_g} câu GHÉP lời gọi + {len(rows) - n_g} câu NGẪU NHIÊN. Đợt 1 ({sum(r['dot'] == 1 for r in rows)} câu, "
         "khoảng 15–20 phút) là bắt buộc; đợt 2, 3 tùy chọn (thêm câu ngẫu nhiên để số liệu chắc hơn).", False),
        ("Hai loại câu", True),
        ("GHÉP LỜI GỌI: máy lấy câu lệnh đã sinh rồi gắn lời gọi trợ lý cùng giọng vào trước. Có câu đủ cụm (\"Trợ lý ơi, …\"), "
         "có câu chỉ còn đuôi (\"…lý ơi, …\") — giả lập bộ phát hiện từ đánh thức cắt không gọn. Mảnh âm tiết bị cắt dở KHÔNG có trong nhãn.", False),
        ("NGẪU NHIÊN: bốc ngẫu nhiên trong 10.000 câu thường, chia đều 3 vùng, không nhìn kết quả chấm của máy. "
         "Dùng để biết tỷ lệ câu đúng thật sự và luật chấm nào (v1 hay v2) khớp tai người.", False),
        ("Cách nghe", True),
        ("1. Sheet \"Nghe\": bấm \"▶ Nghe\" (tệp Excel phải nằm trong thư mục runs/, cạnh ket_qua_full10k). Nghe tối đa 2 lần.", False),
        ("2. Mọi câu: \"Cả câu đúng?\" = Đúng / Sai / Không chắc so với cột \"Nhãn (câu cần nghe)\". Sai thì ghi từ nghe sai.", False),
        ("3. Riêng câu GHÉP, điền thêm:", False),
        ("   • \"Chỗ nối\": Tự nhiên / Có tiếng tách / Lệch âm lượng / Khoảng lặng lạ.", False),
        ("   • \"Nhãn phần gọi\": Khớp (nghe đúng các tiếng trong cột \"Phần gọi theo nhãn\") / Thừa tiếng (nghe rõ thêm một tiếng mà nhãn bỏ) "
         "/ Thiếu tiếng (nhãn ghi mà không nghe thấy).", False),
        ("4. CHƯA mở các cột bị gập bên phải cho tới khi xong cả đợt (đó là máy chấm và máy nghe ra gì).", False),
        ("Thế nào là Đúng", True),
        ("Mọi từ nghe ra đúng là từ trong nhãn. Giọng Trung khác thanh vẫn tính Đúng nếu người địa phương hiểu đúng từ. Từ địa phương tính Đúng.", False),
        ("Giới hạn", True),
        ("Một người nghe, không hoàn toàn mù; 20 câu ghép và 30–100 câu ngẫu nhiên chỉ để định hướng (khoảng tin cậy rộng).", False),
    ]
    for t, bold in lines:
        hd.append([t])
        hd.cell(hd.max_row, 1).font = Font(bold=bold, size=12 if bold else 11)
        hd.cell(hd.max_row, 1).alignment = Alignment(wrap_text=True)
    hd.column_dimensions["A"].width = 130
    hd.append([])
    hd.append(["Tổng hợp (tự cập nhật khi bạn điền)"])
    hd.cell(hd.max_row, 1).font = Font(bold=True, size=12)
    head = ["Nhóm", "Số câu", "Đã nghe", "Đúng", "Sai", "Không chắc", "Nối tự nhiên", "Nhãn gọi khớp"]
    hd.append(head)
    for i in range(1, len(head) + 1):
        hd.cell(hd.max_row, i).font = Font(bold=True)
    for name in ("Ghép lời gọi", "Ngẫu nhiên"):
        q = f'"{name}"'
        hd.append([name, f"=COUNTIF(Nghe!$C:$C,{q})", f'=COUNTIFS(Nghe!$C:$C,{q},Nghe!$G:$G,"<>")',
                   *[f'=COUNTIFS(Nghe!$C:$C,{q},Nghe!$G:$G,"{v}")' for v in ("Đúng", "Sai", "Không chắc")],
                   f'=COUNTIFS(Nghe!$C:$C,{q},Nghe!$I:$I,"Tự nhiên")', f'=COUNTIFS(Nghe!$C:$C,{q},Nghe!$J:$J,"Khớp")'])
    for col in "BCDEFGH":
        hd.column_dimensions[col].width = 12

    ws = wb.create_sheet("Nghe")
    cols = [("STT", 5), ("Đợt", 5), ("Loại", 12), ("Nghe", 8), ("Nhãn (câu cần nghe)", 48), ("Phần gọi theo nhãn", 14),
            ("Cả câu đúng?", 11), ("Từ nghe sai", 22), ("Chỗ nối", 15), ("Nhãn phần gọi", 13), ("Ghi chú", 26),
            ("_", 2), ("_", 2),
            ("Vùng", 8), ("Máy chấm v1", 14), ("Máy chấm v2", 14), ("Kiểu ghép", 18), ("PhoWhisper-large nghe", 40),
            ("PhoWhisper-medium nghe", 40), ("Zipformer nghe", 40), ("PhoWhisper-small (T) nghe", 40), ("Giọng", 12), ("Mã câu", 36)]
    ws.append([c for c, _ in cols])
    for i, (_, w) in enumerate(cols, 1):
        c = ws.cell(1, i)
        c.font = Font(bold=True)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[c.column_letter].width = w
    fill, grey = PatternFill("solid", fgColor="FFF2CC"), PatternFill("solid", fgColor="EEEEEE")
    for k, r in enumerate(rows, 1):
        m, g1, g2 = r["m"], r["g1"], r["g2"]
        is_g = r["nhom"] == "G"
        kieu = ""
        if is_g:
            kieu = "đủ cụm" if m["wake_type"] == "du" else ("đuôi, đúng ranh giới" if m["wake_cat"] == "ranh_gioi"
                                                             else f"đuôi, cắt giữa âm tiết {int(m['wake_cat'][-1]) + 1}")
        ws.append([k, r["dot"], "Ghép lời gọi" if is_g else "Ngẫu nhiên", "▶ Nghe", m["text_raw"],
                   (m.get("wake_giu") or "(không còn tiếng nào)") if is_g else "", "", "", "", "", "", "", "",
                   m.get("voice_region") or m["speaker_region"], VERDICT_VI.get(g1["ket_qua"], g1["ket_qua"]),
                   VERDICT_VI.get(g2["ket_qua"], g2["ket_qua"]), kieu, g1.get("hyp_L", ""), g1.get("hyp_M", ""),
                   g1.get("hyp_Z", ""), g1.get("hyp_T", ""), m["voice"], m["utt_id"]])
        link = ws.cell(k + 1, 4)
        link.hyperlink = f"{REL}/{r['run']}/{m['audio48']}".replace("/", "\\")
        link.font = Font(color="0563C1", underline="single")
        for c in (7, 8, 11):
            ws.cell(k + 1, c).fill = fill
        for c in (9, 10):
            ws.cell(k + 1, c).fill = fill if is_g else grey
        ws.cell(k + 1, 5).alignment = Alignment(wrap_text=True, vertical="top")
    n = len(rows) + 1
    for rng, opts in ((f"G2:G{n}", "Đúng,Sai,Không chắc"), (f"I2:I{n}", "Tự nhiên,Có tiếng tách,Lệch âm lượng,Khoảng lặng lạ"),
                      (f"J2:J{n}", "Khớp,Thừa tiếng,Thiếu tiếng")):
        dv = DataValidation(type="list", formula1=f'"{opts}"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(rng)
    ws.column_dimensions.group("L", "W", hidden=True, outline_level=1)
    ws.freeze_panes = "F2"
    ws.auto_filter.ref = f"A1:W{n}"
    wb.save(path)


def main():
    rows = pick()
    path = os.path.join(HERE, "runs", "nghe_thu_full10k.xlsx")
    if os.path.exists(path):
        path = path.replace(".xlsx", time.strftime("_%Y%m%d_%H%M%S.xlsx"))
    write(path, rows)
    miss = [r for r in rows if not os.path.exists(os.path.join(ROOT, r["run"], r["m"]["audio48"]))]
    print(f"{len(rows)} câu -> {path}")
    print("Theo nhóm × đợt:", dict(collections.Counter((r["nhom"], r["dot"]) for r in rows)))
    print("Câu ghép theo kiểu:", dict(collections.Counter((r["m"]["wake_type"], r["m"]["wake_cat"], r["m"]["wake_phrase"]) for r in rows if r["nhom"] == "G")))
    print("Theo vùng:", dict(collections.Counter((r["nhom"], r["m"]["speaker_region"]) for r in rows)))
    print("Phân bố máy chấm v1 ở câu ngẫu nhiên:", dict(collections.Counter(r["g1"]["ket_qua"] for r in rows if r["nhom"] == "N")))
    print("Thiếu tệp âm thanh:", len(miss))
    with open(path.replace(".xlsx", ".jsonl"), "w", encoding="utf-8", newline="\n") as f:
        for k, r in enumerate(rows, 1):
            f.write(json.dumps(dict(stt=k, dot=r["dot"], nhom=r["nhom"], run=r["run"], utt_id=r["m"]["utt_id"]), ensure_ascii=False) + "\n")


if __name__ == "__main__":
    main()
