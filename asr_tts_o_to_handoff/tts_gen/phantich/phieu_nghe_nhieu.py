# -*- coding: utf-8 -*-
"""Phiếu nghe nhiễu (sau phantich/do_nhieu.py): gán cảnh cho 60 phiên AISHELL-5 và tìm mưa / tiếng người trong
uncontrolled_conditions của CAVEMOVE. Chạy trên laptop, từ tts_gen/:
  python -m phantich.phieu_nghe_nhieu ds --do runs/nhieu_v1/do        # -> <do>/doan_nghe.tsv (đưa lên cụm để cắt)
  (trên cụm) python -m phantich.cat_doan_nghe --ds .../doan_nghe.tsv --nhieu ~/asr_tts_o_to/nhieu --out .../nghe2
  python -m phantich.phieu_nghe_nhieu phieu --do runs/nhieu_v1/do     # -> <do>/nghe2/phieu_nghe_nhieu.xlsx
Bản 2 (29/10, người nghe không rõ đang nghe gì): giữ độ to thật trong mỗi nhóm (cat_doan_nghe), thêm sheet mẫu đối
chiếu từ các tệp CAVEMOVE đã biết nhãn, viết lại hướng dẫn. Nhóm "cavemove" (mẫu + đoạn unc) chung một mức khuếch đại
nên so được độ to; AISHELL-5 là nhóm riêng (xe, micro, mức ghi khác).
Chọn CAVEMOVE unc (tất định theo băm): 12 lộp độp cao tần nhất (mưa?), 12 sáng tiếng nhất (đường ướt? tiếng người?),
4 nhịp chậm mạnh nhất (gạt nước?), 12 ngẫu nhiên làm mốc. Cột số đo bị gập để người nghe không bị dẫn dắt.
"""
import argparse
import csv
import os

from thuvien.chon_cau import h32

A5 = "aishell5/giai_nen/noise/{p}/noise_DX01C01.wav"
CM = "cavemove/giai_nen/cavemove_dataset/{xe}/{ch}/{loai}/{tep}.wav"
# Mẫu đối chiếu: Smart forfour (có đủ đứng yên, đường nhám, quạt 3 mức), micro phân tán kênh 1 (táp-lô).
MAU = [
    ("M1", "noise", "s0_w0", "Xe đứng yên, máy nổ, kính đóng", "Nền rất nhỏ, chỉ có tiếng máy rì rì"),
    ("M2", "noise", "s50_w0_ver1", "Chạy 50 km/h, kính đóng", "Tiếng lốp lăn ù ù đều, trầm"),
    ("M3", "noise", "s110_w0_ver1", "Chạy 110 km/h, kính đóng", "Ù ù to hơn, có tiếng gió rít quanh thân xe"),
    ("M4", "noise", "s60_w2_ver1", "Chạy 60 km/h, mở hết kính trước", "Gió lùa phần phật, ồn tần thấp dao động mạnh"),
    ("M5", "ventilation", "v1_w0", "Đứng yên, quạt điều hòa mức thấp", "Tiếng quạt xì xì nhẹ, đều"),
    ("M6", "ventilation", "v3_w0", "Đứng yên, quạt điều hòa mức cao nhất", "Tiếng quạt xì to, đều, sáng"),
    ("M7", "noise", "s30_w0_coarse", "Chạy 30 km/h trên đường nhám", "Lốp lạo xạo, rung, nhiều tiếng lục cục"),
]
CO_KHONG = "Có,Không,Không chắc"
TRANG_THAI = "Đứng yên,Chạy chậm,Chạy nhanh,Thay đổi trong đoạn,Không rõ"


def doc(path):
    rows = list(csv.DictReader(open(path, encoding="utf-8")))
    for r in rows:
        for k, v in r.items():
            try:
                r[k] = float(v)
            except (TypeError, ValueError):
                pass
    return rows


def chon_cavemove(rows):
    unc = [r for r in rows if r["loai"] == "uncontrolled_conditions"]
    for r in unc:
        r["sang"] = r["b3000_8000"] - r["b300_1000"]
    dung, out = set(), []

    def lay(ds, n, ly_do):
        k = 0
        for r in ds:
            key = (r["xe"], r["cau_hinh"], r["tep"])
            if k < n and key not in dung:
                dung.add(key)
                out.append((ly_do, r))
                k += 1
    lay(sorted(unc, key=lambda r: -r["xung_cao_tan"]), 12, "lộp độp cao tần")
    lay(sorted(unc, key=lambda r: -r["sang"]), 12, "sáng tiếng")
    lay(sorted(unc, key=lambda r: -(r["nhip_r"] or 0)), 4, "nhịp chậm")
    lay(sorted(unc, key=lambda r: h32("phieu_nhieu", r["xe"], r["cau_hinh"], r["tep"])), 12, "ngẫu nhiên (mốc)")
    return sorted(out, key=lambda t: h32("thu_tu", t[1]["xe"], t[1]["cau_hinh"], t[1]["tep"]))


def ten_unc(r):
    return f"{r['xe']}_{r['cau_hinh']}_{r['tep']}"


def ghi_ds(do):
    a5 = sorted(doc(os.path.join(do, "aishell5.csv")), key=lambda r: h32("a5", r["phien"]))
    cm = chon_cavemove(doc(os.path.join(do, "cavemove.csv")))
    dong = ["nhom\tten\ttep\tbat_dau_s\tdai_s"]
    dong += [f"cavemove\t{m}\t{CM.format(xe='Smart_forfour', ch='distributed', loai=l, tep=t)}\t-1\t10"
             for m, l, t, *_ in MAU]
    dong += [f"cavemove\t{ten_unc(r)}\t{CM.format(xe=r['xe'], ch=r['cau_hinh'], loai='uncontrolled_conditions', tep=r['tep'])}"
             "\t0\t0" for _, r in cm]
    dong += [f"aishell5\t{int(r['phien']):03d}\t{A5.format(p=f'{int(r['phien']):03d}')}\t-1\t10" for r in a5]
    p = os.path.join(do, "doan_nghe.tsv")
    with open(p, "w", encoding="utf-8", newline="\n") as f:
        f.write("\n".join(dong) + "\n")
    print(f"{p}: {len(dong) - 1} đoạn ({len(MAU)} mẫu, {len(cm)} CAVEMOVE unc, {len(a5)} AISHELL-5)")


def ghi_phieu(do):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    nghe = os.path.join(do, "nghe2")
    LINK = Font(color="0563C1", underline="single")

    def chon(ws, cot, ds, n):
        dv = DataValidation(type="list", formula1=f'"{ds}"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{cot}2:{cot}{n}")

    def dau_bang(ws, rong):
        for c in ws[1]:
            c.font, c.fill = Font(bold=True), PatternFill("solid", fgColor="DDEBF7")
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for k, w in rong.items():
            ws.column_dimensions[k].width = w
        ws.freeze_panes = "C2"

    wb = Workbook()
    hd = wb.active
    hd.title = "0_Huong_dan"
    for dong in [
        "VÌ SAO CẦN NGHE",
        "Bước tiếp theo trộn tiếng ồn trong xe vào giọng TTS theo 15 tình huống (preset): xe đỗ, chạy phố, cao tốc, mở kính, "
        "quạt mạnh, xe điện chạy chậm... Mỗi tình huống cần đúng loại tiếng ồn.",
        "• AISHELL-5 có 60 bản ghi tiếng ồn trong một xe điện (mỗi bản 30–74 phút) nhưng KHÔNG ghi bản nào là tình huống gì. "
        "Bạn nghe 10 giây giữa mỗi bản và chọn nhãn, để biết dùng bản nào cho preset nào.",
        "• CAVEMOVE có 764 đoạn ngắn 'ghi trong điều kiện tự do', không có nhãn. Cần biết: có đoạn nào MƯA không (nguồn duy "
        "nhất cho preset mưa), và có đoạn nào lẫn TIẾNG NGƯỜI NÓI không (phải loại).",
        "",
        "CÁCH LÀM (khoảng 20–30 phút; không cần nghe hết)",
        "1. Sheet 1_Mau_doi_chieu: nghe 7 mẫu đã biết chắc nhãn để quen tai (đứng yên, chạy chậm, chạy nhanh, mở kính, quạt, "
        "đường nhám).",
        "2. Sheet 2_CAVEMOVE: nghe các đoạn, chọn Có/Không cho mưa, tiếng người, âm báo, gió lùa. Mẫu và các đoạn này cùng "
        "mức khuếch đại, nên to hay nhỏ so với mẫu là thật.",
        "3. Sheet 3_AISHELL5: nghe từng đoạn, chọn xe đứng yên/chạy chậm/chạy nhanh, có gió lùa, nhạc, quạt mạnh, tiếng "
        "người. Độ to giữa các đoạn AISHELL-5 với nhau là thật (đoạn yên sẽ nghe nhỏ); không so độ to với CAVEMOVE (xe "
        "và micro khác).",
        "",
        "LƯU Ý",
        "• Nên dùng tai nghe. Mở phiếu từ đúng thư mục này để liên kết ▶ Nghe chạy được.",
        "• Chỉnh âm lượng một lần cho mẫu M3 (110 km/h) nghe vừa tai rồi giữ nguyên. Khi đó M1 (đứng yên) sẽ rất nhỏ: đó là "
        "khác biệt thật. Cột 'Độ to' cho biết đoạn nhỏ hơn đoạn to nhất của cùng nhóm bao nhiêu dB (0 = to nhất; −30 = "
        "rất nhỏ, có thể phải tăng âm lượng mới nghe rõ có tiếng người hay không).",
        "• Không chắc thì chọn 'Không chắc' hoặc để trống. Ô trống = chưa nghe.",
        "• Phiên 014 và 052 của AISHELL-5 là cùng một bản ghi (sẽ gộp làm một).",
    ]:
        hd.append([dong])
        hd.cell(row=hd.max_row, column=1).alignment = Alignment(wrap_text=True)
        if dong.isupper():
            hd.cell(row=hd.max_row, column=1).font = Font(bold=True)
    hd.column_dimensions["A"].width = 130

    # độ to so với đoạn to nhất của nhóm (cat_doan_nghe đặt đoạn to nhất ở −20 dBFS RMS, trừ khi phải hạ vì đỉnh)
    muc = {(r["nhom"], r["ten"]): float(r["rms_ra_db"])
           for r in csv.DictReader(open(os.path.join(nghe, "muc.tsv"), encoding="utf-8"), delimiter="\t")}
    to_nhat = {g: max(v for (gg, _), v in muc.items() if gg == g) for g, _ in muc}

    def do_to(g, ten):
        return round(muc[(g, ten)] - to_nhat[g], 1)

    wm = wb.create_sheet("1_Mau_doi_chieu")
    wm.append(["Mã", "▶ Nghe (10 s)", "Độ to (dB)", "Tình huống (đã biết)", "Nghe như thế nào"])
    for m, loai, tep, nhan, mo_ta in MAU:
        wm.append([m, "▶ Nghe", do_to("cavemove", m), nhan, mo_ta])
        c = wm.cell(row=wm.max_row, column=2)
        c.hyperlink, c.font = f"cavemove/{m}.flac", LINK
    dau_bang(wm, {"A": 6, "B": 16, "C": 10, "D": 42, "E": 60})

    wc = wb.create_sheet("2_CAVEMOVE")
    wc.append(["Mã", "▶ Nghe (5–21 s)", "Độ to (dB)", "Mưa / đường ướt (lộp độp, xì xèo nước)", "Có tiếng người nói",
               "Âm báo / xi-nhan (tích tắc, bíp)", "Gió lùa (kính mở)", "Ghi chú",
               "xe", "cấu hình", "tệp", "lý do chọn", "dài (s)", "mức dBFS", "lộp độp", "sáng tiếng (dB)"])
    for i, (ly_do, r) in enumerate(chon_cavemove(doc(os.path.join(do, "cavemove.csv"))), 1):
        wc.append([f"C{i:02d}", "▶ Nghe", do_to("cavemove", ten_unc(r)), "", "", "", "", "", r["xe"], r["cau_hinh"],
                   r["tep"], ly_do, r["dai_s"], r["muc_db"], r["xung_cao_tan"], round(r["sang"], 1)])
        c = wc.cell(row=wc.max_row, column=2)
        c.hyperlink, c.font = f"cavemove/{ten_unc(r)}.flac", LINK
    for cot in "DEFG":
        chon(wc, cot, CO_KHONG, wc.max_row)
    for k in "IJKLMNOP":
        wc.column_dimensions[k].hidden = True
    dau_bang(wc, {"A": 6, "B": 16, "C": 10, "D": 20, "E": 14, "F": 18, "G": 14, "H": 30})

    wa = wb.create_sheet("3_AISHELL5")
    wa.append(["Mã", "▶ Nghe (10 s giữa phiên)", "Độ to (dB)", "Xe đang làm gì", "Gió lùa (kính mở)", "Có nhạc",
               "Quạt gió mạnh", "Có tiếng người nói", "Ghi chú",
               "phiên", "số kênh", "dài (phút)", "mức dBFS", "dải 300 Hz–1 kHz", "dao động p90−p10"])
    for i, r in enumerate(sorted(doc(os.path.join(do, "aishell5.csv")), key=lambda r: h32("a5", r["phien"])), 1):
        p = f"{int(r['phien']):03d}"
        wa.append([f"A{i:02d}", "▶ Nghe", do_to("aishell5", p), "", "", "", "", "", "", p, int(r["so_kenh"]),
                   round(r["dai_s"] / 60, 1), r["muc_db"], r["b300_1000"], r["dao_dong_db"]])
        c = wa.cell(row=wa.max_row, column=2)
        c.hyperlink, c.font = f"aishell5/{p}.flac", LINK
    chon(wa, "D", TRANG_THAI, wa.max_row)
    for cot in "EFGH":
        chon(wa, cot, CO_KHONG, wa.max_row)
    for k in "JKLMNO":
        wa.column_dimensions[k].hidden = True
    dau_bang(wa, {"A": 6, "B": 20, "C": 10, "D": 20, "E": 14, "F": 10, "G": 12, "H": 14, "I": 30})

    out = os.path.join(nghe, "phieu_nghe_nhieu.xlsx")
    wb.save(out)
    thieu = [c.hyperlink.target for s in (wm, wc, wa) for c in s["B"][1:]
             if not os.path.exists(os.path.join(nghe, c.hyperlink.target))]
    print(f"Phiếu: {out} ({wm.max_row - 1} mẫu, {wc.max_row - 1} CAVEMOVE, {wa.max_row - 1} AISHELL-5); liên kết thiếu: {thieu[:3] or 'không'}")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("viec", choices=["ds", "phieu"])
    ap.add_argument("--do", required=True)
    a = ap.parse_args()
    (ghi_ds if a.viec == "ds" else ghi_phieu)(a.do)


if __name__ == "__main__":
    main()
