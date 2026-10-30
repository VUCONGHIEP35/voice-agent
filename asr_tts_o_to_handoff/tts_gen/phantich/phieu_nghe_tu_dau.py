# -*- coding: utf-8 -*-
"""Phiếu nghe thử các câu MỞ ĐẦU bằng lời gọi / từ đệm ("Này", "Ờ", "Ê", "Nè", "Trợ lý…") trong 1.000 câu đánh giá.

Chạy (từ tts_gen/, venv có openpyxl):  .venv_zip/Scripts/python -m phantich.phieu_nghe_tu_dau
Ghi runs/nghe_thu_tu_dau.xlsx (không ghi đè: trùng tên thì thêm giờ). Tệp phải nằm trong runs/ vì nút "▶ Nghe" là liên kết
tương đối tới runs/ket_qua_A/... và runs/ket_qua_B/... (FLAC 48 kHz).
Chọn câu: dòng hiện hành, tts_check đạt, >= 3 từ (như phantich/vi_tri_loi.py), từ đầu thuộc nhóm lời gọi/đệm hoặc "trợ lý";
thêm 12 câu đối chứng mở đầu bằng từ nội dung mà L hoặc M nghe sai từ đầu. Nhóm:
  A lời gọi/đệm, L hoặc M sai từ đầu    B lời gọi/đệm, L và M đúng từ đầu
  C "Trợ lý…", L hoặc M sai từ đầu      D "Trợ lý…", L và M đúng từ đầu
  E đối chứng: từ nội dung, L hoặc M sai từ đầu
Chia đợt khoảng 30 câu, đợt nào cũng trộn đủ nhóm; thứ tự trong đợt xáo theo seed cố định.
Các cột máy (nhóm, kết quả chấm, câu L/M/Z/T nghe ra) bị gập lại để người nghe không bị dẫn dắt.
"""
import collections
import csv
import json
import os
import random
import time

from cauhinh.duong_dan import GOC as HERE  # noqa: E402
from thuvien.cham_chu import norm  # noqa: E402
from thuvien.van_ban import load_lexicon  # noqa: E402
from thuvien.manifest import active_rows, load_manifest  # noqa: E402
from thuvien.can_hang import word_states  # noqa: E402

RUNS = [  # (thư mục manifest + audio48, thư mục chứa gate_hop_server.csv, đường dẫn tương đối từ runs/)
    ("runs/ket_qua_A/runs/eval1000_A", "runs/ket_qua_B/runs/ket_qua_A/runs/eval1000_A", "ket_qua_A/runs/eval1000_A"),
    ("runs/ket_qua_B/runs/eval1000_B", "runs/ket_qua_B/runs/eval1000_B", "ket_qua_B/runs/eval1000_B"),
]
FILL = {"này", "ờ", "ê", "ơ", "ừ", "à", "nè", "nì", "ủa"}
GROUPS = {"A": "A · Lời gọi/đệm · máy sai từ đầu", "B": "B · Lời gọi/đệm · máy đúng từ đầu",
          "C": "C · \"Trợ lý…\" · máy sai từ đầu", "D": "D · \"Trợ lý…\" · máy đúng từ đầu",
          "E": "E · Đối chứng: từ nội dung · máy sai từ đầu"}
QUESTIONS = {
    "A": "Âm thanh từ đầu có thật sự sai/nhẹ quá, hay chỉ máy nghe nhầm?",
    "B": "Đối chứng trong nhóm: máy đúng thì tai người cũng thấy rõ chứ?",
    "C": "\"Trợ lý\" đọc có rõ không, hay bị nuốt/đọc sai thanh?",
    "D": "Đối chứng: \"Trợ lý\" máy nghe đúng, tai người thấy thế nào?",
    "E": "Câu bình thường mà máy sai từ đầu: do TTS hay do ASR?",
}
SEED = 20261025
N_CONTROL = 12
BATCH = 30


def pick():
    lex = load_lexicon(os.path.join(HERE, "lexicon", "asr_bien_the.tsv"))
    lex.sort(key=lambda x: -len(x[0]))
    out, control = [], []
    for run, gate_dir, rel in RUNS:
        man = {m["utt_id"]: m for m in active_rows(load_manifest(os.path.join(HERE, run)))}
        with open(os.path.join(HERE, gate_dir, "gate_hop_server.csv"), encoding="utf-8-sig") as f:
            gate = list(csv.DictReader(f))
        for g in gate:
            m = man.get(g["utt_id"])
            if not m or g["tts_check"] != "dat":
                continue
            words = norm(m["text_spoken"]).split()
            if len(words) < 3:
                continue
            st = {c: word_states(m, g.get(f"hyp_{c}", ""), lex) for c in ("L", "M")}
            wrong = any(s and s[0] != "ok" for s in st.values())
            kind = "fill" if words[0] in FILL else ("tro_ly" if words[:2] == ["trợ", "lý"] else "noi_dung")
            row = dict(m=m, g=g, audio=f"{rel}/{m['audio48']}", wrong=wrong)
            if kind == "fill":
                out.append(dict(row, nhom="A" if wrong else "B"))
            elif kind == "tro_ly":
                out.append(dict(row, nhom="C" if wrong else "D"))
            elif wrong:
                control.append(dict(row, nhom="E"))
    rng = random.Random(SEED)
    out += rng.sample(sorted(control, key=lambda r: r["m"]["utt_id"]), min(N_CONTROL, len(control)))
    # chia đợt: rải từng nhóm đều qua các đợt rồi xáo trong đợt
    by = collections.defaultdict(list)
    for r in sorted(out, key=lambda r: r["m"]["utt_id"]):
        by[r["nhom"]].append(r)
    n_batch = max(1, round(len(out) / BATCH))
    batches = [[] for _ in range(n_batch)]
    for k in sorted(by):
        rng.shuffle(by[k])
        for i, r in enumerate(by[k]):
            batches[i % n_batch].append(r)
    res = []
    for b, rows in enumerate(batches, 1):
        rng.shuffle(rows)
        for r in rows:
            res.append(dict(r, dot=b))
    return res


def first_heard(hyp):
    w = norm(hyp or "").split()
    return " ".join(w[:2])


def write(path, rows):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = Workbook()
    hd = wb.active
    hd.title = "Hướng dẫn"
    n_dot = max(r["dot"] for r in rows)
    lines = [
        ("Nghe thử: câu mở đầu bằng lời gọi / từ đệm (1.000 câu đánh giá A+B)", True),
        (f"{len(rows)} câu, chia {n_dot} đợt khoảng {BATCH} câu (mỗi đợt 10–15 phút). Làm xong đợt 1 rồi mới làm đợt 2.", False),
        ("Vì sao có phiếu này", True),
        ("Đo trên 1.000 câu: từ ĐẦU câu bị máy nghe sai gấp ~1,6 lần từ giữa câu (gần như không bị MẤT, mà bị chép sai), "
         "nhất là khi câu mở đầu bằng \"Này\", \"Ờ\", \"Ê\", \"Trợ lý\". Cần tai người để biết lỗi do TTS đọc hay do máy nghe.", False),
        ("Cách nghe", True),
        ("1. Ở sheet \"Nghe\", bấm \"▶ Nghe\" để mở tệp âm thanh (tệp Excel phải nằm trong thư mục runs/, cạnh ket_qua_A và ket_qua_B).", False),
        ("2. Đọc câu ở cột \"Câu cần đọc\", nghe tối đa 2 lần. Chọn \"Cả câu đúng?\" (Đúng / Sai / Không chắc).", False),
        ("3. Cột \"Từ đầu\" (QUAN TRỌNG): Rõ / Nhẹ-ngắn (nghe ra nhưng yếu, lướt) / Sai âm (thành tiếng khác) / Mất (không nghe thấy). "
         "Nếu không Rõ, ghi tiếng bạn nghe ra ở cột kế bên.", False),
        ("4. CHƯA mở các cột bị gập ở bên phải (bấm dấu + phía trên cột) cho tới khi xong cả đợt: các cột đó là máy chấm "
         "và máy nghe ra gì, xem trước sẽ nghe theo máy.", False),
        ("Thế nào là Đúng", True),
        ("Mọi từ nghe ra đúng là từ trong câu. Giọng Trung khác thanh vẫn tính Đúng nếu người địa phương hiểu đúng từ. "
         "Từ địa phương (tui, nha, nì) tính Đúng.", False),
        ("Giới hạn", True),
        ("Một người nghe, không hoàn toàn mù (bạn biết dự án): kết quả để định hướng. Câu Trung: nếu không phải người miền Trung, "
         "phân vân thì chọn Không chắc.", False),
        ("Câu hỏi mỗi nhóm trả lời", True),
    ] + [(f"{GROUPS[k]}: {QUESTIONS[k]}", False) for k in GROUPS]
    for t, bold in lines:
        hd.append([t])
        hd.cell(hd.max_row, 1).font = Font(bold=bold, size=12 if bold else 11)
        hd.cell(hd.max_row, 1).alignment = Alignment(wrap_text=True)
    hd.column_dimensions["A"].width = 130
    hd.append([])
    hd.append(["Tổng hợp (tự cập nhật khi bạn điền)"])
    hd.cell(hd.max_row, 1).font = Font(bold=True, size=12)
    head = ["Nhóm", "Số câu", "Đã nghe", "Cả câu Đúng", "Sai", "Không chắc", "Từ đầu Rõ", "Nhẹ-ngắn", "Sai âm", "Mất"]
    hd.append(head)
    for i in range(1, len(head) + 1):
        hd.cell(hd.max_row, i).font = Font(bold=True)
    for k, name in GROUPS.items():
        q = f'"{name}"'
        rng_g, rng_f, rng_h = "Nghe!$N:$N", "Nghe!$F:$F", "Nghe!$G:$G"
        hd.append([name, f"=COUNTIF({rng_g},{q})", f'=COUNTIFS({rng_g},{q},{rng_f},"<>")',
                   *[f'=COUNTIFS({rng_g},{q},{rng_f},"{v}")' for v in ("Đúng", "Sai", "Không chắc")],
                   *[f'=COUNTIFS({rng_g},{q},{rng_h},"{v}")' for v in ("Rõ", "Nhẹ-ngắn", "Sai âm", "Mất")]])
    for col in "BCDEFGHIJ":
        hd.column_dimensions[col].width = 11

    ws = wb.create_sheet("Nghe")
    cols = [("STT", 5), ("Đợt", 5), ("Nghe", 8), ("Câu cần đọc", 46), ("Vùng giọng", 8), ("Cả câu đúng?", 11),
            ("Từ đầu", 11), ("Từ đầu nghe ra (nếu không Rõ)", 18), ("Từ khác nghe sai", 22), ("Ghi chú", 26),
            # các cột máy, gập lại
            ("_", 2), ("_", 2), ("_", 2), ("Nhóm", 30), ("Máy chấm", 14), ("Từ đầu: L nghe", 14), ("Từ đầu: M nghe", 14),
            ("PhoWhisper-large nghe", 40), ("PhoWhisper-medium nghe", 40), ("Zipformer nghe", 40), ("PhoWhisper-small (T) nghe", 40),
            ("Giọng", 12), ("Mã câu", 34)]
    ws.append([c for c, _ in cols])
    for i, (_, w) in enumerate(cols, 1):
        c = ws.cell(1, i)
        c.font = Font(bold=True)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[c.column_letter].width = w
    fill = PatternFill("solid", fgColor="FFF2CC")
    for k, r in enumerate(rows, 1):
        m, g = r["m"], r["g"]
        ws.append([k, r["dot"], "▶ Nghe", m["text_spoken"], g["region"], "", "", "", "", "", "", "", "",
                   GROUPS[r["nhom"]], {"dat": "Đạt", "dat_1nguon": "Đạt một nguồn", "can_nghe": "Cần nghe", "loai": "Loại"}.get(
                       g["ket_qua"], g["ket_qua"]), first_heard(g.get("hyp_L")), first_heard(g.get("hyp_M")),
                   g.get("hyp_L", ""), g.get("hyp_M", ""), g.get("hyp_Z", ""), g.get("hyp_T", ""), m["voice"], m["utt_id"]])
        link = ws.cell(k + 1, 3)
        link.hyperlink = r["audio"].replace("/", "\\")
        link.font = Font(color="0563C1", underline="single")
        for c in range(6, 11):
            ws.cell(k + 1, c).fill = fill
        ws.cell(k + 1, 4).alignment = Alignment(wrap_text=True, vertical="top")
    n = len(rows) + 1
    dv1 = DataValidation(type="list", formula1='"Đúng,Sai,Không chắc"', allow_blank=True)
    dv2 = DataValidation(type="list", formula1='"Rõ,Nhẹ-ngắn,Sai âm,Mất"', allow_blank=True)
    ws.add_data_validation(dv1)
    ws.add_data_validation(dv2)
    dv1.add(f"F2:F{n}")
    dv2.add(f"G2:G{n}")
    ws.column_dimensions.group("K", "W", hidden=True, outline_level=1)
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = f"A1:W{n}"
    wb.save(path)


def main():
    rows = pick()
    path = os.path.join(HERE, "runs", "nghe_thu_tu_dau.xlsx")
    if os.path.exists(path):
        path = path.replace(".xlsx", time.strftime("_%Y%m%d_%H%M%S.xlsx"))
    write(path, rows)
    missing = [r["audio"] for r in rows if not os.path.exists(os.path.join(HERE, "runs", r["audio"]))]
    print(f"{len(rows)} câu, {max(r['dot'] for r in rows)} đợt -> {path}")
    print("Theo nhóm:", dict(collections.Counter(r["nhom"] for r in rows)))
    print("Theo đợt × nhóm:", {d: dict(collections.Counter(r["nhom"] for r in rows if r["dot"] == d))
                              for d in sorted({r["dot"] for r in rows})})
    print("Theo vùng:", dict(collections.Counter(r["g"]["region"] for r in rows)))
    print("Tệp âm thanh không tìm thấy:", missing[:5] or "không")
    with open(path.replace(".xlsx", ".jsonl"), "w", encoding="utf-8", newline="\n") as f:
        for k, r in enumerate(rows, 1):
            f.write(json.dumps(dict(stt=k, dot=r["dot"], nhom=r["nhom"], utt_id=r["m"]["utt_id"], audio=r["audio"]),
                               ensure_ascii=False) + "\n")


if __name__ == "__main__":
    main()
