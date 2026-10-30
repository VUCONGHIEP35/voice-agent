# -*- coding: utf-8 -*-
"""Kiểm và nhận dạng tệp giọng thật (WAV/MP3) trong một thư mục bằng Zipformer-30M; ghi kết quả ra cùng thư mục.

Chạy (từ tts_gen/, venv riêng):  .venv_zip/Scripts/python -m buoc.giong_that giong_that
Ghi vào thư mục đó (không ghi đè; trùng tên thì thêm giờ chạy):
  ket_qua_zipformer.xlsx  sheet "Kết quả" (mỗi tệp một dòng: thông số âm thanh, cảnh báo, câu Z nghe được, hai cột trống
                          để người dùng điền câu thực đã đọc và Z đúng/sai) + sheet "Tổng hợp"
  ket_qua_zipformer.jsonl cùng nội dung, cho script khác đọc
Tệp coi như dữ liệu chưa tin cậy: chỉ nhận .wav/.mp3 <= 20 MB, kiểm byte đầu (RIFF..WAVE / ID3 / khung MPEG), đọc bằng
soundfile; tệp lỗi được ghi lý do, không làm dừng cả lượt. Không sửa, không chép tệp âm thanh đi đâu.
Mô hình: Z = Zipformer-30M int8, modified_beam_search beam 4 (như cấu hình Z của buoc/cham.py).
"""
import argparse
import json
import os
import time

from buoc.asr_zipformer import MODEL_DIR, SR16, make_tag, model_files  # noqa: E402

EXTS = (".wav", ".mp3")
MAX_BYTES = 20 * 2 ** 20
MIN_S, WARN_S, MAX_S = 0.3, 15.0, 120.0


def magic_ok(path):
    with open(path, "rb") as f:
        b = f.read(12)
    if path.lower().endswith(".wav"):
        return b[:4] == b"RIFF" and b[8:12] == b"WAVE"
    return b[:3] == b"ID3" or (len(b) >= 2 and b[0] == 0xFF and (b[1] & 0xE0) == 0xE0)


def db(x):
    import math
    return round(20 * math.log10(max(float(x), 1e-10)), 1)


def measure(w, sr):
    """Thông số đơn giản trên tín hiệu mono float32."""
    import numpy as np
    peak = float(np.abs(w).max()) if len(w) else 0.0
    clip = float(np.mean(np.abs(w) >= 0.999)) if len(w) else 0.0
    head = w[: max(1, int(0.005 * sr))]
    head_db = db(np.sqrt(np.mean(head ** 2)))
    hop = int(0.02 * sr)
    frames = [w[i:i + hop] for i in range(0, len(w) - hop + 1, hop)]
    rms = np.array([np.sqrt(np.mean(f ** 2)) for f in frames]) if frames else np.array([0.0])
    # Ngưỡng "có tiếng": cao hơn nền ồn (phân vị 10 của RMS khung 20 ms) 10 dB, tối thiểu -40 dBFS. Dùng -40 dBFS cố định
    # thì tệp thu điện thoại có nền khoảng -33 dBFS sẽ bị coi là "không có im lặng" dù có.
    noise = float(np.percentile(rms, 10))
    thr = max(10 ** (-40 / 20), noise * 10 ** (10 / 20))
    loud = np.where(rms > thr)[0]
    lead = round(loud[0] * hop / sr, 2) if len(loud) else None
    trail = round((len(rms) - 1 - loud[-1]) * hop / sr, 2) if len(loud) else None
    snr = db(np.percentile(rms, 90) / max(noise, 1e-10))
    return dict(dinh_dbfs=db(peak), xen_ty_le=round(clip, 5), dau_5ms_dbfs=head_db, nen_on_dbfs=db(noise),
                im_lang_dau_s=lead, im_lang_cuoi_s=trail, snr_uoc_db=snr)


def warnings_for(info, m, dur):
    w = []
    if info["samplerate"] < 16000:
        w.append(f"tần số lấy mẫu {info['samplerate']} Hz < 16 kHz")
    if info["channels"] > 1:
        w.append(f"{info['channels']} kênh, đã lấy trung bình")
    if dur > WARN_S:
        w.append(f"dài {dur:.1f} s (> {WARN_S:.0f} s, quy cách 1 câu/tệp)")
    if m["xen_ty_le"] > 0.001:
        w.append(f"xén đỉnh {m['xen_ty_le']:.2%}")
    if m["dinh_dbfs"] < -30:
        w.append(f"thu nhỏ (đỉnh {m['dinh_dbfs']} dBFS)")
    if m["dau_5ms_dbfs"] > -40:
        w.append("không có im lặng ở đầu (5 ms đầu > -40 dBFS)")
    elif m["im_lang_dau_s"] is not None and m["im_lang_dau_s"] < 0.2:
        w.append(f"im lặng đầu chỉ {m['im_lang_dau_s']} s")
    if m["snr_uoc_db"] < 15:
        w.append(f"ồn (SNR ước {m['snr_uoc_db']} dB)")
    return w


def load(path):
    import numpy as np
    import soundfile as sf
    import soxr
    info = sf.info(path)
    w, sr = sf.read(path, dtype="float32", always_2d=True)
    w = w.mean(axis=1)
    w16 = soxr.resample(w, sr, SR16) if sr != SR16 and len(w) else w
    return dict(format=info.format, subtype=info.subtype, samplerate=info.samplerate, channels=info.channels), w, sr, \
        np.ascontiguousarray(w16, dtype=np.float32)


def out_path(folder, name):
    p = os.path.join(folder, name)
    if not os.path.exists(p):
        return p
    base, ext = os.path.splitext(name)
    return os.path.join(folder, f"{base}_{time.strftime('%Y%m%d_%H%M%S')}{ext}")


def write_xlsx(path, rows, summary):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Kết quả"
    cols = [("STT", 5), ("Tệp", 28), ("Z nghe được", 60), ("Câu thực đã đọc (bạn điền)", 45), ("Z đúng? (Đúng/Sai/Gần đúng)", 16),
            ("Cảnh báo", 45), ("Thời lượng (s)", 10), ("Định dạng", 14), ("Tần số (Hz)", 10), ("Kênh", 6), ("Đỉnh (dBFS)", 10),
            ("Xén đỉnh", 9), ("5 ms đầu (dBFS)", 11), ("Nền ồn (dBFS)", 10), ("Im lặng đầu (s)", 10), ("Im lặng cuối (s)", 10), ("SNR ước (dB)", 10),
            ("Lỗi đọc tệp", 30)]
    ws.append([c for c, _ in cols])
    for i, (_, wdt) in enumerate(cols, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = wdt
        ws.cell(1, i).font = Font(bold=True)
        ws.cell(1, i).alignment = Alignment(wrap_text=True, vertical="top")
    fill = PatternFill("solid", fgColor="FFF2CC")
    for k, r in enumerate(rows, 1):
        ws.append([k, r["tep"], r.get("z_nghe", ""), "", "", "; ".join(r.get("canh_bao", [])), r.get("thoi_luong_s"),
                   f"{r.get('format', '')}/{r.get('subtype', '')}" if r.get("format") else "", r.get("samplerate"),
                   r.get("channels"), r.get("dinh_dbfs"), r.get("xen_ty_le"), r.get("dau_5ms_dbfs"), r.get("nen_on_dbfs"),
                   r.get("im_lang_dau_s"),
                   r.get("im_lang_cuoi_s"), r.get("snr_uoc_db"), r.get("loi", "")])
        for c in (4, 5):
            ws.cell(k + 1, c).fill = fill
        for c in (3, 4, 6):
            ws.cell(k + 1, c).alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "C2"
    s = wb.create_sheet("Tổng hợp")
    for key, val in summary:
        s.append([key, val])
    s.column_dimensions["A"].width = 34
    s.column_dimensions["B"].width = 90
    for r in s.iter_rows(min_col=1, max_col=1):
        r[0].font = Font(bold=True)
    wb.save(path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("folder")
    ap.add_argument("--model-dir", default=MODEL_DIR)
    ap.add_argument("--threads", type=int, default=4)
    a = ap.parse_args()
    import sherpa_onnx
    files = sorted(f for f in os.listdir(a.folder) if os.path.isfile(os.path.join(a.folder, f)))
    audio = [f for f in files if f.lower().endswith(EXTS)]
    skipped = [f for f in files if not f.lower().endswith(EXTS) and not f.startswith("ket_qua_zipformer")]
    f = model_files(a.model_dir, "int8")
    t0 = time.time()
    rec = sherpa_onnx.OfflineRecognizer.from_transducer(
        encoder=f["encoder"], decoder=f["decoder"], joiner=f["joiner"], tokens=f["tokens"], num_threads=a.threads,
        sample_rate=SR16, feature_dim=80, decoding_method="modified_beam_search", max_active_paths=4, provider="cpu")
    load_s = time.time() - t0
    rows, dec_s, audio_s = [], 0.0, 0.0
    for name in audio:
        p = os.path.join(a.folder, name)
        r = dict(tep=name, byte=os.path.getsize(p))
        if r["byte"] > MAX_BYTES:
            r["loi"] = f"lớn hơn {MAX_BYTES // 2 ** 20} MB"
        elif r["byte"] == 0:
            r["loi"] = "tệp rỗng"
        elif not magic_ok(p):
            r["loi"] = "byte đầu không đúng định dạng của đuôi tệp"
        if "loi" not in r:
            try:
                info, w, sr, w16 = load(p)
            except Exception as e:  # tệp hỏng: ghi lý do, đi tiếp
                r["loi"] = f"không đọc được: {type(e).__name__}: {e}"
            else:
                dur = len(w) / sr
                r.update(info, thoi_luong_s=round(dur, 2), **measure(w, sr))
                r["canh_bao"] = warnings_for(info, r, dur)
                if dur < MIN_S or dur > MAX_S:
                    r["loi"] = f"thời lượng {dur:.2f} s ngoài khoảng {MIN_S}–{MAX_S} s"
                else:
                    t1 = time.time()
                    s = rec.create_stream()
                    s.accept_waveform(SR16, w16)
                    rec.decode_stream(s)
                    r["z_nghe"] = s.result.text.strip()
                    dec_s += time.time() - t1
                    audio_s += dur
        rows.append(r)
        print(f"{name}: {r.get('thoi_luong_s', '?')} s | {r.get('z_nghe', '')} | {'; '.join(r.get('canh_bao', []))} {r.get('loi', '')}")
    tag = make_tag("int8", "modified_beam_search", 4)
    ok = [r for r in rows if "z_nghe" in r]
    summary = [("Thư mục", os.path.basename(os.path.abspath(a.folder))), ("Lúc chạy", time.strftime("%Y-%m-%d %H:%M:%S")),
               ("Mô hình", f"hynt/Zipformer-30M-RNNT-6000h ({tag}), sherpa-onnx {sherpa_onnx.__version__}, CPU {a.threads} luồng"),
               ("Giấy phép mô hình", "CC BY-NC-ND 4.0 (chỉ đánh giá, phi thương mại)"),
               ("Tệp âm thanh", len(audio)), ("Nhận dạng được", len(ok)), ("Lỗi", len(rows) - len(ok)),
               ("Tệp có cảnh báo", sum(bool(r.get("canh_bao")) for r in rows)),
               ("Tổng thời lượng (s)", round(audio_s, 1)),
               ("Thời gian nạp mô hình / nhận dạng (s)", f"{load_s:.1f} / {dec_s:.1f}"),
               ("Tệp khác bị bỏ qua", ", ".join(skipped) or "không"),
               ("Chưa chấm CER", "Tên tệp không có mã câu (sample_id) nên chưa biết câu đáp án; điền cột "
                                "'Câu thực đã đọc' để chấm. Quy cách đặt tên: <sample_id>__<dieu_kien>.wav"),
               ("Chữ HOA", "Z xuất chữ hoa, không dấu câu; khi chấm sẽ chuẩn hóa như asr_check.norm")]
    for r in rows:
        summary.append((r["tep"], r.get("z_nghe") or r.get("loi", "")))
    xp = out_path(a.folder, "ket_qua_zipformer.xlsx")
    write_xlsx(xp, rows, summary)
    jp = out_path(a.folder, "ket_qua_zipformer.jsonl")
    with open(jp, "w", encoding="utf-8", newline="\n") as fo:
        for r in rows:
            fo.write(json.dumps(dict(r, model=tag), ensure_ascii=False) + "\n")
    print(f"Ghi {xp}\nGhi {jp}")


if __name__ == "__main__":
    main()
